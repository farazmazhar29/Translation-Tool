import io
from typing import Any

from conftest import Headers, upload
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

TEXT = "Der Korpus und die Frontblende sind mit Art. AB-1234 für die Küche. Preis 49,90 EUR."


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Montageanleitung", level=1)
    doc.add_paragraph(TEXT)
    doc.add_paragraph("12345")  # no letters -> not translatable
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"], ws["B1"], ws["C1"] = "Der Griff ist weiß", 5, "=B1*2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_login_and_permissions(client: TestClient, translator: Headers) -> None:
    assert client.post("/api/auth/login", data={"username": "admin", "password": "x"}).status_code == 401
    assert client.get("/api/projects").status_code == 401
    assert client.get("/api/users", headers=translator).status_code == 403
    assert client.get("/api/auth/me", headers=translator).json()["role"] == "translator"


def test_upload_rejects_unsupported(client: TestClient, translator: Headers) -> None:
    r = client.post("/api/projects", headers=translator, files={"file": ("a.imdd", b"x")})
    assert r.status_code == 400
    assert "InDesign" in r.json()["detail"]
    assert client.post("/api/projects", headers=translator, files={"file": ("a.docx", b"broken")}).status_code == 400


def test_full_workflow(client: TestClient, translator: Headers, reviewer: Headers) -> None:
    p = upload(client, translator, "anleitung.docx", make_docx(), name="Test", customer="ACME", target_langs="en")
    assert p["detected_lang"] == "de"
    assert p["target_langs"] == ["en"]
    assert p["status"] == "new"
    pid = p["id"]

    # preview only -> export must be blocked until approval
    assert client.post(f"/api/projects/{pid}/translate", headers=translator).status_code == 200
    detail = client.get(f"/api/projects/{pid}", headers=translator).json()
    assert detail["status"] == "preview_ready"
    assert len(detail["segments"]) == 2
    seg = detail["segments"][1]
    tr = seg["translations"]["en"]
    assert "carcass" in tr["text"]
    assert "front panel" in tr["text"]
    assert {"glossary", "protected"} <= {t["category"] for t in tr["terms"]}
    assert any(t["source"] == "AB-1234" and t["dnt"] for t in tr["terms"])
    body = {"lang": "en", "format": "docx"}
    assert client.post(f"/api/projects/{pid}/export", headers=translator, json=body).status_code == 409

    # manual edit + comment, quality check, translators cannot approve
    r = client.put(
        f"/api/segments/{seg['id']}/en", headers=translator, json={"text": "Edited text", "comment": "check"}
    )
    assert r.json()["translations"]["en"]["edited"] is True
    q = client.get(f"/api/projects/{pid}/quality", headers=translator).json()
    assert any(w["type"] == "glossary_violation" for w in q["warnings"])  # the edit removed the glossary terms
    assert client.post(f"/api/projects/{pid}/approve", headers=translator).status_code == 403

    # regenerate the segment -> glossary terms are back, then approve as reviewer
    assert client.post(f"/api/segments/{seg['id']}/en/retranslate", headers=translator).status_code == 200
    assert client.post(f"/api/projects/{pid}/approve", headers=reviewer).json()["status"] == "approved"
    assert client.put(f"/api/segments/{seg['id']}/en", headers=translator, json={"text": "x"}).status_code == 409

    # exports in all formats
    for fmt in ("docx", "xlsx", "pdf"):
        for mode in ("translation", "both"):
            opts = {"include_report": True, "include_comments": True, "archive_after": False}
            r = client.post(
                f"/api/projects/{pid}/export", headers=translator, json={**body, "format": fmt, "mode": mode, **opts}
            )
            assert r.status_code == 200, r.text
            assert client.get(f"/api/exports/{r.json()['id']}/download", headers=translator).content
    hist = client.get(f"/api/projects/{pid}/history", headers=translator).json()
    assert len(hist["releases"]) == 1
    assert len(hist["exports"]) == 6
    assert hist["releases"][0]["approved_by"] == "reviewer"

    # archive, search, copy, reopen
    client.post(f"/api/projects/{pid}/archive", headers=translator)
    assert pid not in [x["id"] for x in client.get("/api/projects", headers=translator).json()]
    params = {"archived": "only", "term": "carcass", "customer": "acme"}
    assert [x["id"] for x in client.get("/api/projects", params=params, headers=translator).json()] == [pid]
    copy = client.post(f"/api/projects/{pid}/copy", headers=translator, params={"with_translations": True}).json()
    assert copy["status"] == "preview_ready"
    assert copy["name"].endswith("(Kopie)")
    assert client.post(f"/api/projects/{pid}/reopen", headers=translator).json()["status"] == "preview_ready"


def test_xlsx_project_and_custom_columns(client: TestClient, translator: Headers, reviewer: Headers) -> None:
    pid = upload(client, translator, "preise.xlsx", make_xlsx(), target_langs="en")["id"]
    client.post(f"/api/projects/{pid}/translate", headers=translator)
    client.post(f"/api/projects/{pid}/approve", headers=reviewer)
    cfg = {"source_col": "B", "target_col": "D", "start_row": 3, "sheet_name": "Export"}
    body: dict[str, Any] = {"lang": "en", "format": "xlsx", "archive_after": False}
    r = client.post(f"/api/projects/{pid}/export", headers=translator, json={**body, "mode": "both", "config": cfg})
    assert r.status_code == 200
    content = client.get(f"/api/exports/{r.json()['id']}/download", headers=translator).content
    ws = load_workbook(io.BytesIO(content))["Export"]
    assert ws["B3"].value == "Der Griff ist weiß"
    assert ws["D3"].value == "The Handle is white"
    assert ws["D2"].value
    # translation-only keeps the original workbook (formula and number untouched)
    r = client.post(f"/api/projects/{pid}/export", headers=translator, json=body)
    orig = load_workbook(
        io.BytesIO(client.get(f"/api/exports/{r.json()['id']}/download", headers=translator).content)
    ).active
    assert orig is not None
    assert orig["A1"].value == "The Handle is white"
    assert orig["C1"].value == "=B1*2"


def test_glossary(client: TestClient, admin: Headers, translator: Headers) -> None:
    new = {"terms": {"de": "Wangenblende", "en": "end panel"}, "synonyms": {"en": ["side panel"]}}
    r = client.post("/api/glossary", headers=translator, json=new)
    assert r.status_code == 200
    assert r.json()["status"] == "draft"
    assert client.post("/api/glossary", headers=translator, json=new).status_code == 409  # duplicate check
    eid = r.json()["id"]
    assert client.put(f"/api/glossary/{eid}", headers=translator, json=new).status_code == 403
    r = client.put(f"/api/glossary/{eid}", headers=admin, json={**new, "status": "approved", "description": "v2"})
    assert r.json()["version"] == 2
    assert [v["action"] for v in client.get(f"/api/glossary/{eid}/versions", headers=admin).json()] == [
        "update",
        "create",
    ]
    dnt = {"terms": {"de": "ProLine X"}, "do_not_translate": True, "category": "product"}
    assert client.post("/api/glossary", headers=admin, json=dnt).json()["status"] == "approved"
    assert client.post("/api/glossary", headers=admin, json={"terms": {"de": "nur deutsch"}}).status_code == 400

    wb = load_workbook(io.BytesIO(client.get("/api/glossary/export/xlsx", headers=admin).content))
    ws = wb.active
    assert ws is not None
    ws.append(["Fronttür", "door", "", "", "", "", "technical", "approved"])
    ws.append(["Wangenblende", "end panel"])  # duplicate
    ws.append(["kaputt"])  # invalid
    buf = io.BytesIO()
    wb.save(buf)
    res = client.post("/api/glossary/import/xlsx", headers=admin, files={"file": ("g.xlsx", buf.getvalue())}).json()
    reasons = [r["reason"] for r in res["skipped_rows"]]  # re-imported export rows are duplicates too
    assert res["created"] == 1
    assert "Begriff oder Übersetzung fehlt" in reasons
    assert "Dublette: Wangenblende" in reasons
    assert client.delete(f"/api/glossary/{eid}", headers=translator).status_code == 403
    assert client.delete(f"/api/glossary/{eid}", headers=admin).status_code == 200


def test_dashboard_and_settings(client: TestClient, admin: Headers, translator: Headers) -> None:
    d = client.get("/api/dashboard", headers=translator).json()
    assert d["stats"]["projects_total"] >= 2
    assert d["top_terms"]
    assert client.put("/api/settings", headers=translator, json={}).status_code == 403
    s = client.put("/api/settings", headers=admin, json={"default_export_format": "pdf", "evil": 1}).json()
    assert s["default_export_format"] == "pdf"
