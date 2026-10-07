"""Glossary management (CRUD, duplicates, versions, Excel) and the admin endpoints (settings, dashboard)."""

import io
import uuid
from typing import Any

import pytest
from conftest import Headers, docx_bytes, upload
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.config import settings

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def word() -> str:
    return "t" + uuid.uuid4().hex[:8]  # unique, so tests never collide with entries of other tests


def entry(de: str, en: str, **over: Any) -> dict[str, Any]:
    return {"terms": {"de": de, "en": en}} | over


def create(client: TestClient, headers: Headers, body: dict[str, Any]) -> dict[str, Any]:
    r = client.post("/api/glossary", headers=headers, json=body)
    assert r.status_code == 200, r.text
    return r.json()


def sheet_bytes(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_xlsx(client: TestClient, headers: Headers, rows: list[list[Any]], **params: Any) -> dict[str, Any]:
    r = client.post(
        "/api/glossary/import/xlsx", headers=headers, files={"file": ("g.xlsx", sheet_bytes(rows))}, params=params
    )
    assert r.status_code == 200, r.text
    return r.json()


# Creating entries
def test_translators_propose_drafts_reviewers_publish(
    client: TestClient, translator: Headers, reviewer: Headers
) -> None:
    proposed = create(client, translator, entry(word(), word(), status="approved"))  # a translator cannot self-approve
    assert proposed["status"] == "draft"
    assert proposed["updated_by"] == "translator"
    assert create(client, reviewer, entry(word(), word()))["status"] == "approved"
    assert create(client, reviewer, entry(word(), word(), status="draft"))["status"] == "draft"


@pytest.mark.parametrize(
    "body",
    [
        {"terms": {}},  # nothing
        {"terms": {"fr": "x", "es": "y"}},  # unsupported languages
        {"terms": {"de": "nur deutsch"}},  # a translation is required ...
        {"terms": {"de": "  ", "en": "  "}},  # ... and blanks do not count
    ],
)
def test_incomplete_entries_are_rejected(client: TestClient, translator: Headers, body: dict[str, Any]) -> None:
    assert client.post("/api/glossary", headers=translator, json=body).status_code == 400


def test_do_not_translate_needs_only_one_language(client: TestClient, reviewer: Headers) -> None:
    name = word()
    created = create(client, reviewer, {"terms": {"de": name}, "do_not_translate": True, "category": "product"})
    assert created["do_not_translate"] is True
    assert created["terms"] == {"de": name}


def test_terms_are_trimmed_and_synonyms_cleaned(client: TestClient, reviewer: Headers) -> None:
    a, b = word(), word()
    created = create(
        client, reviewer, entry(f"  {a}  ", f" {b} ", synonyms={"de": [" x1 ", "", "x2"], "fr": ["ignored"]})
    )
    assert created["terms"] == {"de": a, "en": b}
    assert created["synonyms"] == {"de": ["x1", "x2"]}


@pytest.mark.parametrize(
    "body",
    [{"terms": {"de": "x", "en": "y"}, "category": "weird"}, {"terms": {"de": "x", "en": "y"}, "status": "bogus"}],
)
def test_invalid_category_or_status_is_a_422(client: TestClient, reviewer: Headers, body: dict[str, Any]) -> None:
    assert client.post("/api/glossary", headers=reviewer, json=body).status_code == 422


# Duplicates
def test_duplicates_are_detected_case_and_space_insensitively(client: TestClient, reviewer: Headers) -> None:
    a = word()
    original = create(client, reviewer, entry(f"{a} Korpus", f"{a} carcass", synonyms={"de": [f"{a}-Kasten"]}))
    for clash in (
        entry(f"  {a.upper()}   korpus ", word()),
        entry(word(), f"{a} CARCASS"),
        entry(f"{a}-kasten", word()),
    ):
        r = client.post("/api/glossary", headers=reviewer, json=clash)
        assert r.status_code == 409
        assert original["id"] in [d["id"] for d in r.json()["detail"]["duplicates"]]
    forced = client.post("/api/glossary", headers=reviewer, json=entry(f"{a} korpus", word(), force=True))
    assert forced.status_code == 200  # the user may insist


def test_updating_an_entry_does_not_clash_with_itself_but_with_others(client: TestClient, reviewer: Headers) -> None:
    a, b = create(client, reviewer, entry(word(), word())), create(client, reviewer, entry(word(), word()))
    same = {"terms": a["terms"], "description": "neu"}
    assert client.put(f"/api/glossary/{a['id']}", headers=reviewer, json=same).status_code == 200
    clash = {"terms": {"de": b["terms"]["de"], "en": word()}}
    assert client.put(f"/api/glossary/{a['id']}", headers=reviewer, json=clash).status_code == 409


def test_duplicate_scan_lists_pairs(client: TestClient, reviewer: Headers) -> None:
    a = word()
    first = create(client, reviewer, entry(a, word()))
    second = create(client, reviewer, entry(a, word(), force=True))
    groups = [g for g in client.get("/api/glossary/duplicates/scan", headers=reviewer).json() if g["term"] == a]
    assert groups == [{"term": a, "lang": "de", "ids": [first["id"], second["id"]]}]


# Update, status, delete, versions
def test_update_bumps_the_version_and_keeps_the_history(
    client: TestClient, reviewer: Headers, translator: Headers, admin: Headers
) -> None:
    created = create(client, reviewer, entry(word(), word()))
    eid = created["id"]
    assert client.put(f"/api/glossary/{eid}", headers=translator, json=entry("x", "y")).status_code == 403
    changed = client.put(f"/api/glossary/{eid}", headers=reviewer, json=entry(word(), word(), description="v2")).json()
    assert changed["version"] == 2
    assert (
        client.post(f"/api/glossary/{eid}/status", headers=reviewer, params={"status": "rejected"}).json()["version"]
        == 3
    )
    history = client.get(f"/api/glossary/{eid}/versions", headers=translator).json()
    assert [(v["entry_version"], v["action"]) for v in history] == [(3, "update"), (2, "update"), (1, "create")]
    assert history[0]["data"]["status"] == "rejected"
    assert history[2]["data"]["terms"] == created["terms"]

    assert client.delete(f"/api/glossary/{eid}", headers=reviewer).status_code == 403  # admins delete
    assert client.delete(f"/api/glossary/{eid}", headers=admin).status_code == 200
    after = client.get(f"/api/glossary/{eid}/versions", headers=admin).json()
    assert after[0]["action"] == "delete"  # the audit trail outlives the entry
    assert eid not in [e["id"] for e in client.get("/api/glossary", headers=admin).json()]


def test_status_workflow(client: TestClient, translator: Headers, reviewer: Headers) -> None:
    draft = create(client, translator, entry(word(), word()))
    url = f"/api/glossary/{draft['id']}/status"
    assert client.post(url, headers=translator, params={"status": "approved"}).status_code == 403
    assert (
        client.post(url, headers=reviewer, params={"status": "bogus"}).status_code == 422
    )  # validated against the enum
    assert client.post(url, headers=reviewer, params={"status": "approved"}).json()["status"] == "approved"
    assert (
        client.post("/api/glossary/999999/status", headers=reviewer, params={"status": "approved"}).status_code == 404
    )


def test_unknown_entries_are_404(client: TestClient, reviewer: Headers, admin: Headers) -> None:
    assert client.put("/api/glossary/999999", headers=reviewer, json=entry("a", "b")).status_code == 404
    assert client.delete("/api/glossary/999999", headers=admin).status_code == 404


def test_rejected_entries_are_not_used_by_the_translation(
    client: TestClient, reviewer: Headers, translator: Headers
) -> None:
    term = word()
    created = create(client, reviewer, entry(term, f"{term}-english"))
    client.post(f"/api/glossary/{created['id']}/status", headers=reviewer, params={"status": "rejected"})
    pid = upload(client, translator, "g.docx", docx_bytes(f"Der {term} ist weiß und gut."))["id"]
    client.post(f"/api/projects/{pid}/translate", headers=translator)
    terms = client.get(f"/api/projects/{pid}", headers=translator).json()["segments"][0]["translations"]["en"]["terms"]
    assert term not in [t["source"] for t in terms if t["category"] == "glossary"]


# Searching and listing
def test_list_filters_and_limit(client: TestClient, reviewer: Headers, translator: Headers) -> None:
    t = word()
    create(client, reviewer, entry(f"{t} Schrank", f"{t} cabinet", description=f"Möbel {t}"))
    create(client, reviewer, entry(f"{t} Griff", f"{t} handle", synonyms={"de": [f"{t}-Knauf"]}, category="product"))
    create(client, translator, entry(f"{t} Tür", f"{t} door"))

    def lookup(**params: str) -> set[str]:
        rows = client.get("/api/glossary", headers=translator, params={"q": t} | params).json()
        return {e["terms"]["de"] for e in rows}

    assert len(lookup()) == 3
    assert lookup(status="draft") == {f"{t} Tür"}
    assert lookup(category="product") == {f"{t} Griff"}
    assert lookup(q=f"{t}-knauf") == {f"{t} Griff"}  # synonyms are searchable
    assert lookup(q=f"möbel {t}") == {f"{t} Schrank"}  # so are descriptions
    assert len(client.get("/api/glossary", headers=translator, params={"limit": 1}).json()) == 1
    assert client.get("/api/glossary", headers=translator, params={"limit": 0}).status_code == 422


# Excel export / import
def test_excel_export_has_the_documented_columns(client: TestClient, reviewer: Headers) -> None:
    t = word()
    create(client, reviewer, entry(f"{t} a", f"{t} b", synonyms={"de": ["s1", "s2"]}, description="Beschreibung"))
    r = client.get("/api/glossary/export/xlsx", headers=reviewer)
    assert r.headers["content-type"] == XLSX
    assert "glossar.xlsx" in r.headers["content-disposition"]
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))  # pyright: ignore[reportOptionalMemberAccess]
    assert rows[0] == (
        "de",
        "en",
        "de_synonyme",
        "en_synonyme",
        "beschreibung",
        "nicht_uebersetzen",
        "kategorie",
        "status",
    )
    mine = next(row for row in rows if row[0] == f"{t} a")
    assert mine[1:5] == (f"{t} b", "s1; s2", None, "Beschreibung")


def test_excel_import_creates_updates_and_reports_what_it_skipped(client: TestClient, reviewer: Headers) -> None:
    t = word()
    existing = create(client, reviewer, entry(f"{t} alt", f"{t} old"))
    header = ["de", "en", "de_synonyme", "en_synonyme", "beschreibung", "nicht_uebersetzen", "kategorie", "status"]
    rows = [
        header,
        [
            f"{t} neu",
            f"{t} new",
            "n1; n2|n3",
            "",
            "Neu",
            "",
            "technical",
            "draft",
        ],  # created, synonyms split on ; and |
        [
            f"{t} marke",
            None,
            None,
            None,
            "Marke",
            "ja",
            "product",
            "approved",
        ],  # do-not-translate needs one language only
        [f"{t} alt", f"{t} other", None, None, None, None, None, None],  # duplicate of an existing entry
        [f"{t} halb", None, None, None, None, None, None, None],  # translation missing
        [None, None, None, None, None, None, None, None],  # empty line
        [
            f"{t} komisch",
            f"{t} weird",
            None,
            None,
            None,
            None,
            "irgendwas",
            "bogus",
        ],  # invalid category/status fall back
    ]
    result = import_xlsx(client, reviewer, rows)
    assert (result["created"], result["updated"], result["skipped"]) == (3, 0, 3)
    assert {(r["row"], r["reason"].split(":")[0]) for r in result["skipped_rows"]} == {
        (4, "Dublette"),
        (5, "Begriff oder Übersetzung fehlt"),
        (6, "Begriff oder Übersetzung fehlt"),
    }
    # the codes let the interface word the reasons in its own language; a duplicate names its term
    assert {(r["row"], r["code"], r["term"] != "") for r in result["skipped_rows"]} == {
        (4, "duplicate", True),
        (5, "missing_term", False),
        (6, "missing_term", False),
    }

    stored = {e["terms"]["de"]: e for e in client.get("/api/glossary", headers=reviewer, params={"q": t}).json()}
    assert stored[f"{t} neu"]["synonyms"]["de"] == ["n1", "n2", "n3"]
    assert stored[f"{t} neu"]["status"] == "draft"
    assert stored[f"{t} marke"]["do_not_translate"] is True
    assert (stored[f"{t} komisch"]["category"], stored[f"{t} komisch"]["status"]) == ("technical", "approved")

    update = import_xlsx(
        client,
        reviewer,
        [header, [f"{t} alt", f"{t} replaced", "alias", None, "aktualisiert", None, None, None]],
        update_existing=True,
    )
    assert (update["created"], update["updated"]) == (0, 1)
    changed = client.get("/api/glossary", headers=reviewer, params={"q": f"{t} alt"}).json()[0]
    assert (changed["id"], changed["version"], changed["description"]) == (existing["id"], 2, "aktualisiert")
    assert changed["terms"]["en"] == f"{t} replaced"


def test_excel_import_understands_column_aliases(client: TestClient, reviewer: Headers) -> None:
    t = word()
    result = import_xlsx(
        client, reviewer, [["Deutsch", "Englisch", "Description", "DNT"], [f"{t} x", f"{t} y", "Alias-Test", None]]
    )
    assert result["created"] == 1
    assert client.get("/api/glossary", headers=reviewer, params={"q": t}).json()[0]["description"] == "Alias-Test"


@pytest.mark.parametrize(
    ("payload", "message"),
    [
        (b"this is not a spreadsheet", "gelesen"),
        (sheet_bytes([]), "Kopfzeile"),
        (sheet_bytes([["foo", "bar"], ["1", "2"]]), "Kopfzeile"),
    ],
)
def test_bad_excel_files_are_rejected(client: TestClient, reviewer: Headers, payload: bytes, message: str) -> None:
    r = client.post("/api/glossary/import/xlsx", headers=reviewer, files={"file": ("g.xlsx", payload)})
    assert r.status_code == 400
    assert message in r.json()["detail"]


def test_only_reviewers_import(client: TestClient, translator: Headers) -> None:
    r = client.post(
        "/api/glossary/import/xlsx", headers=translator, files={"file": ("g.xlsx", sheet_bytes([["de", "en"]]))}
    )
    assert r.status_code == 403


# Admin endpoints: settings and dashboard
def test_settings_defaults_merge_and_ignore_unknown_keys(
    client: TestClient, admin: Headers, translator: Headers
) -> None:
    before = client.get("/api/settings", headers=translator).json()
    assert before["gemini_key_configured"] is False
    assert set(before["colors"]) == {"technical", "glossary", "ambiguous", "unknown", "protected"}
    changed = client.put(
        "/api/settings",
        headers=admin,
        json={"default_export_format": "pdf", "colors": {**before["colors"], "glossary": "#000000"}, "hack": 1},
    ).json()
    assert changed["default_export_format"] == "pdf"
    assert changed["colors"]["glossary"] == "#000000"
    assert "hack" not in changed
    assert changed["llm_model"] == before["llm_model"]  # untouched values are kept
    client.put("/api/settings", headers=admin, json={"default_export_format": "docx", "colors": before["colors"]})
    assert client.put("/api/settings", headers=translator, json={}).status_code == 403


def test_the_model_is_chosen_from_a_list_not_typed(client: TestClient, admin: Headers) -> None:
    choices = client.get("/api/settings", headers=admin).json()["model_choices"]
    # the configured model first, then its fallbacks
    assert choices["gemini"] == list(dict.fromkeys([settings.gemini_model, *settings.gemini_fallback_models]))
    assert choices["mock"] == []  # the offline mock needs no model
    original = client.get("/api/settings", headers=admin).json()
    try:
        typo = client.put("/api/settings", headers=admin, json={"llm_provider": "gemini", "llm_model": "gemini-flsh"})
        assert typo.status_code == 400
        assert settings.gemini_model in typo.json()["detail"]  # says what is possible
        assert client.put("/api/settings", headers=admin, json={"llm_provider": "openai"}).status_code == 400
        other = choices["gemini"][-1]
        saved = client.put("/api/settings", headers=admin, json={"llm_provider": "gemini", "llm_model": other})
        assert saved.json()["llm_model"] == other
    finally:
        client.put(
            "/api/settings",
            headers=admin,
            json={"llm_provider": original["llm_provider"], "llm_model": original["llm_model"]},
        )


def test_dashboard_reports_work_in_progress(
    client: TestClient, translator: Headers, reviewer: Headers, translated_project: int
) -> None:
    client.post(f"/api/projects/{translated_project}/approve", headers=reviewer)
    d = client.get("/api/dashboard", headers=translator).json()
    assert {"open", "in_review", "approved", "edited_today", "recent", "top_terms", "stats"} <= d.keys()
    assert d["approved"] >= 1
    assert len(d["recent"]) <= 6
    assert len(d["stats"]["releases_per_day"]) == 14
    assert sum(d["stats"]["releases_per_day"].values()) >= 1
    counts = [t["count"] for t in d["top_terms"]]
    assert counts == sorted(counts, reverse=True)
    assert d["stats"]["glossary_entries"] > 0
    assert d["stats"]["by_language"]["en"] > 0
