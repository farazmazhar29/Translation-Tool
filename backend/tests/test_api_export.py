"""Export: every format and mode, the Excel layout, templates, and the safety of stored files."""

import io
from pathlib import Path
from typing import Any
from urllib.parse import quote

import pymupdf
import pytest
from conftest import Headers, docx_bytes, upload
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.core.config import settings
from app.core.db import SessionLocal
from app.core.enums import QualityCode, Severity
from app.export.models import ExportFile
from app.export.service import number_text, report_lines
from app.quality.schemas import QualityWarning

FORMATS = ("docx", "xlsx", "pdf")


@pytest.fixture
def approved(client: TestClient, reviewer: Headers, translated_project: int) -> int:
    assert client.post(f"/api/projects/{translated_project}/approve", headers=reviewer).status_code == 200
    return translated_project


def export(client: TestClient, headers: Headers, pid: int, **over: Any) -> dict[str, Any]:
    body = {"lang": "en", "format": "docx", "mode": "translation", "archive_after": False} | over
    r = client.post(f"/api/projects/{pid}/export", headers=headers, json=body)
    assert r.status_code == 200, r.text
    return r.json()


def download(client: TestClient, headers: Headers, export_id: int) -> bytes:
    r = client.get(f"/api/exports/{export_id}/download", headers=headers)
    assert r.status_code == 200
    return r.content


def sheet(content: bytes, name: str | None = None) -> Worksheet:
    wb = load_workbook(io.BytesIO(content))
    ws = wb[name] if name else wb.active
    assert isinstance(ws, Worksheet)
    return ws


def pdf_text(content: bytes) -> str:
    with pymupdf.open(stream=content, filetype="pdf") as pdf:
        return "".join(page.get_text("text") for page in pdf)


# Every format and mode produces a valid file
@pytest.mark.parametrize("fmt", FORMATS)
@pytest.mark.parametrize("mode", ["translation", "both"])
def test_every_format_and_mode_yields_a_readable_file_with_the_translation(
    client: TestClient, translator: Headers, approved: int, fmt: str, mode: str
) -> None:
    result = export(client, translator, approved, format=fmt, mode=mode, include_report=True, include_comments=True)
    content = download(client, translator, result["id"])
    if fmt == "docx":
        doc = Document(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
            c.text for t in doc.tables for r in t.rows for c in r.cells
        )
    elif fmt == "xlsx":
        ws = sheet(content)
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    else:
        assert content.startswith(b"%PDF-")
        text = pdf_text(content)
    assert "carcass" in text  # the translated term made it into the file
    assert ("Der Korpus" in text) == (mode == "both")  # the original only when asked for


def test_export_records_who_what_and_when_in_the_history(
    client: TestClient, translator: Headers, approved: int
) -> None:
    export(client, translator, approved, format="pdf", mode="both")
    (entry, *_) = client.get(f"/api/projects/{approved}/history", headers=translator).json()["exports"]
    assert (entry["format"], entry["lang"], entry["mode"], entry["created_by"]) == ("pdf", "en", "both", "translator")
    assert entry["filename"].endswith("_en_both.pdf")


# Word: formatting survives
def test_word_export_keeps_the_original_formatting(client: TestClient, translator: Headers, reviewer: Headers) -> None:
    pid = upload(
        client,
        translator,
        "fmt.docx",
        docx_bytes("Der Korpus ist weiß und die Tür ist neu.", bold_first=True),
    )["id"]
    client.post(f"/api/projects/{pid}/translate", headers=translator)
    client.post(f"/api/projects/{pid}/approve", headers=reviewer)
    doc = Document(io.BytesIO(download(client, translator, export(client, translator, pid)["id"])))
    assert doc.paragraphs[0].text.startswith("The carcass is white")
    assert doc.paragraphs[0].runs[0].bold is True


# An Excel file stays a table in Word and PDF
def parts_list() -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "Ersatzteilliste"
    ws["A3"], ws["B3"], ws["C3"] = "Art.-Nr.", "Bezeichnung", "Preis"
    ws["A4"], ws["B4"], ws["C4"] = "SC-110", "Der Korpus ist weiß und gut.", 12.9
    ws["C4"].number_format = "#,##0.00"
    for ref in ("A1", "A3", "B3", "C3"):
        ws[ref].font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def approved_sheet(client: TestClient, translator: Headers, reviewer: Headers) -> int:
    pid: int = upload(client, translator, "teile.xlsx", parts_list())["id"]
    client.post(f"/api/projects/{pid}/translate", headers=translator)
    assert client.post(f"/api/projects/{pid}/approve", headers=reviewer).status_code == 200
    return pid


def test_an_excel_file_exported_to_word_is_a_table_of_its_rows_and_columns(
    client: TestClient, translator: Headers, approved_sheet: int
) -> None:
    doc = Document(io.BytesIO(download(client, translator, export(client, translator, approved_sheet)["id"])))
    (table,) = doc.tables
    lines = [[c.text for c in row.cells] for row in table.rows]
    assert len(lines) == 3  # title, column heads, one part; the empty row 2 is left out
    assert len(set(lines[0])) == 1  # the title spans the table (merged cells repeat their text)
    assert lines[1] == ["Art.-Nr.", "Bezeichnung", "Preis"] or "Price" in lines[1][2]
    item, name, price = lines[2]
    assert (item, price) == ("SC-110", "12.90")  # a number keeps the decimals of its cell format
    assert "carcass" in name
    assert table.rows[2].cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_original_and_translation_of_an_excel_file_share_each_cell(
    client: TestClient, translator: Headers, approved_sheet: int
) -> None:
    content = download(client, translator, export(client, translator, approved_sheet, mode="both")["id"])
    cell = Document(io.BytesIO(content)).tables[0].rows[2].cells[1]
    assert [p.text for p in cell.paragraphs][1] == "Der Korpus ist weiß und gut."  # the original below


def test_an_excel_file_exported_to_pdf_is_a_table_too(
    client: TestClient, translator: Headers, approved_sheet: int
) -> None:
    content = download(client, translator, export(client, translator, approved_sheet, format="pdf")["id"])
    with pymupdf.open(stream=content, filetype="pdf") as pdf:  # a table row is one text block, its cells as lines
        rows = [text.split("\n") for page in pdf for *_, text, _, _ in page.get_text("blocks")]
    title = next(i for i, row in enumerate(rows) if row[0] == "Ersatzteilliste")
    part = next(i for i, row in enumerate(rows) if row[0] == "SC-110")
    assert "12.90" in rows[part]  # the part and its price on one row, as in the sheet
    assert part > title


@pytest.mark.parametrize(("lang", "expected"), [("de", "1.234,50"), ("en", "1,234.50")])
def test_numbers_are_written_as_the_reader_of_the_translation_writes_them(lang: str, expected: str) -> None:
    assert number_text(1234.5, 2, lang) == expected


def test_word_side_by_side_table_has_the_expected_columns(
    client: TestClient, translator: Headers, approved: int
) -> None:
    content = download(
        client, translator, export(client, translator, approved, mode="both", include_comments=True)["id"]
    )
    header = [c.text for c in Document(io.BytesIO(content)).tables[0].rows[0].cells]
    assert header == ["#", "Original", "Translation", "Comment"]  # an English translation: English words


def test_report_and_comments_are_included_on_request(
    client: TestClient, translator: Headers, reviewer: Headers, translated_project: int
) -> None:
    sid = client.get(f"/api/projects/{translated_project}", headers=translator).json()["segments"][0]["id"]
    client.put(f"/api/segments/{sid}/en", headers=translator, json={"comment": "Prüfen: Maße"})
    client.post(f"/api/projects/{translated_project}/approve", headers=reviewer)
    plain = Document(io.BytesIO(download(client, translator, export(client, translator, translated_project)["id"])))
    rich = Document(
        io.BytesIO(
            download(
                client,
                translator,
                export(client, translator, translated_project, include_report=True, include_comments=True)["id"],
            )
        )
    )
    text = "\n".join(p.text for p in rich.paragraphs)
    assert "Quality report" in text
    assert "Comment: Prüfen: Maße" in text  # the comment itself is what the user wrote
    assert "Quality report" not in "\n".join(p.text for p in plain.paragraphs)


# Excel layout ("Erweiterte Einstellungen")
def test_excel_layout_defaults_and_custom_columns(client: TestClient, translator: Headers, approved: int) -> None:
    default = sheet(
        download(client, translator, export(client, translator, approved, format="xlsx", mode="both")["id"])
    )
    assert default.title == "Translation"
    assert [c.value for c in default[1]][:2] == ["Original", "Translation (EN)"]  # header row
    assert default["A2"].value == "Montageanleitung"

    layout = {
        "sheet_name": "Export",
        "source_col": "c",
        "target_col": "F",
        "comment_col": "",
        "location_col": "H",
        "start_row": 4,
        "header": True,
    }
    custom = sheet(
        download(
            client, translator, export(client, translator, approved, format="xlsx", mode="both", config=layout)["id"]
        ),
        "Export",
    )
    assert custom["C3"].value == "Original"  # the header sits right above the first data row
    assert custom["C4"].value == "Montageanleitung"
    assert custom["F4"].value == "Montageanleitung"  # the mock leaves unknown words unchanged
    assert custom["H4"].value == "Paragraph 1"
    assert custom["A4"].value is None
    assert custom["B4"].value is None


def test_excel_layout_without_header_row(client: TestClient, translator: Headers, approved: int) -> None:
    ws = sheet(
        download(
            client,
            translator,
            export(client, translator, approved, format="xlsx", config={"header": False, "start_row": 1})["id"],
        )
    )
    assert ws["B1"].value  # data starts in the first row


def test_excel_report_goes_to_a_separate_sheet(client: TestClient, translator: Headers, approved: int) -> None:
    wb = load_workbook(
        io.BytesIO(
            download(client, translator, export(client, translator, approved, format="xlsx", include_report=True)["id"])
        )
    )
    assert wb.sheetnames[-1] == "Quality report"


@pytest.mark.parametrize(
    "layout",
    [
        {"target_col": "1"},  # not a column
        {"target_col": ""},  # required
        {"source_col": "AAAA"},  # too long
        {"target_col": "ZZZ"},  # beyond Excel's last column (XFD)
        {"source_col": "B", "target_col": "b"},  # the same column twice would overwrite data
        {"start_row": 0},
        {"sheet_name": ""},
        {"sheet_name": "x" * 32},  # Excel's limit is 31
        {"sheet_name": "bad/name"},
    ],
)
def test_invalid_excel_layouts_are_rejected_with_422(
    client: TestClient, translator: Headers, approved: int, layout: dict[str, Any]
) -> None:
    body = {"lang": "en", "format": "xlsx", "config": layout}
    assert client.post(f"/api/projects/{approved}/export", headers=translator, json=body).status_code == 422


@pytest.mark.parametrize(
    "body", [{"format": "odt"}, {"mode": "everything"}, {"lang": "en", "format": "docx", "template_id": "x"}]
)
def test_invalid_export_requests_are_rejected(
    client: TestClient, translator: Headers, approved: int, body: dict[str, Any]
) -> None:
    assert (
        client.post(f"/api/projects/{approved}/export", headers=translator, json={"lang": "en"} | body).status_code
        == 422
    )


# Templates
def test_export_template_lifecycle_and_use(
    client: TestClient, translator: Headers, reviewer: Headers, approved: int
) -> None:
    layout = {
        "sheet_name": "Vorlage",
        "source_col": "A",
        "target_col": "D",
        "comment_col": "E",
        "location_col": "",
        "start_row": 3,
        "header": True,
    }
    created = client.post(
        "/api/export-templates", headers=translator, json={"name": "Kundenformat", "format": "xlsx", "config": layout}
    )
    assert created.status_code == 200
    tid = created.json()["id"]
    assert created.json()["config"]["target_col"] == "D"
    assert "Kundenformat" in [t["name"] for t in client.get("/api/export-templates", headers=translator).json()]

    ws = sheet(
        download(
            client, translator, export(client, translator, approved, format="xlsx", mode="both", template_id=tid)["id"]
        ),
        "Vorlage",
    )
    assert ws["A2"].value == "Original"  # header row sits right above the first data row (start_row 3)
    assert ws["D3"].value  # the template's target column

    assert client.delete(f"/api/export-templates/{tid}", headers=translator).status_code == 403  # reviewers delete
    assert client.delete(f"/api/export-templates/{tid}", headers=reviewer).status_code == 200
    assert tid not in [t["id"] for t in client.get("/api/export-templates", headers=translator).json()]
    assert client.delete(f"/api/export-templates/{tid}", headers=reviewer).status_code == 200  # idempotent


def test_templates_must_have_a_valid_layout(client: TestClient, translator: Headers) -> None:
    bad = {"name": "kaputt", "format": "xlsx", "config": {"target_col": "1"}}
    assert client.post("/api/export-templates", headers=translator, json=bad).status_code == 422
    assert (
        client.post(
            "/api/export-templates", headers=translator, json={"name": "x", "format": "docx", "config": {}}
        ).status_code
        == 422
    )


# Archiving after export, and the safety of stored files
def test_export_archives_the_project_when_asked(client: TestClient, translator: Headers, approved: int) -> None:
    approved_before = client.get("/api/dashboard", headers=translator).json()["approved"]
    export(client, translator, approved, archive_after=True)
    project = client.get(f"/api/projects/{approved}", headers=translator).json()
    assert (project["status"], project["archived"]) == ("archived", True)
    # archived work is still approved work on the dashboard
    assert client.get("/api/dashboard", headers=translator).json()["approved"] == approved_before
    export(client, translator, approved, archive_after=False)  # exporting again from an archived project is allowed


def test_same_filename_in_two_projects_never_overwrites_an_export(
    client: TestClient, translator: Headers, reviewer: Headers
) -> None:
    """Regression: exports used to be stored as <filename>_<lang>_<mode>.<ext>, so equal filenames collided."""
    ids: list[int] = []
    for text in ("Der Korpus ist weiß und gut.", "Die Frontblende ist neu und gut."):
        pid = upload(client, translator, "plan.docx", docx_bytes(text))["id"]
        client.post(f"/api/projects/{pid}/translate", headers=translator)
        client.post(f"/api/projects/{pid}/approve", headers=reviewer)
        ids.append(export(client, translator, pid)["id"])
    first, second = (Document(io.BytesIO(download(client, translator, i))).paragraphs[0].text for i in ids)
    assert "carcass" in first
    assert "front panel" in second
    with SessionLocal() as db:
        paths = [db.get(ExportFile, i).path for i in ids]  # pyright: ignore[reportOptionalMemberAccess]
    assert len(set(paths)) == 2
    assert all(Path(p).parent == settings.data_dir / "exports" for p in paths)


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("Mein Plan (v2).docx", "Mein_Plan_v2_en_translation.docx"),
        ("../../evil.docx", "evil_en_translation.docx"),
        ("äöü?*.docx", "äöü_en_translation.docx"),
    ],
)
def test_download_names_are_sanitised_user_input(
    client: TestClient, translator: Headers, reviewer: Headers, filename: str, expected: str
) -> None:
    pid = upload(client, translator, filename, docx_bytes("Der Korpus ist weiß und gut."))["id"]
    client.post(f"/api/projects/{pid}/translate", headers=translator)
    client.post(f"/api/projects/{pid}/approve", headers=reviewer)
    assert export(client, translator, pid)["filename"] == expected
    r = client.get(f"/api/exports/{export(client, translator, pid)['id']}/download", headers=translator)
    assert (
        quote(expected) in r.headers["content-disposition"] or expected in r.headers["content-disposition"]
    )  # RFC 5987 for non-ASCII


def test_downloading_unknown_or_deleted_exports_is_404(client: TestClient, translator: Headers, approved: int) -> None:
    assert client.get("/api/exports/999999/download", headers=translator).status_code == 404
    result = export(client, translator, approved)
    with SessionLocal() as db:
        Path(db.get(ExportFile, result["id"]).path).unlink()  # pyright: ignore[reportOptionalMemberAccess]
    assert client.get(f"/api/exports/{result['id']}/download", headers=translator).status_code == 404


def test_export_of_an_unknown_project_and_unauthenticated_access(client: TestClient, translator: Headers) -> None:
    assert client.post("/api/projects/999999/export", headers=translator, json={"lang": "en"}).status_code == 404
    assert client.get("/api/exports/1/download").status_code == 401


# The fixed words follow the language of the translation
def test_a_german_translation_gets_german_headings_and_a_german_report(
    client: TestClient, translator: Headers, reviewer: Headers
) -> None:
    body = "The carcass is white. Price 49.90 EUR."
    pid: int = upload(client, translator, "en.docx", docx_bytes("Assembly", body))["id"]
    client.post(f"/api/projects/{pid}/translate", headers=translator)
    client.post(f"/api/projects/{pid}/approve", headers=reviewer)
    content = download(
        client,
        translator,
        export(client, translator, pid, lang="de", format="xlsx", mode="both", include_report=True)["id"],
    )
    wb = load_workbook(io.BytesIO(content))
    ws = wb.worksheets[0]
    assert ws.title == "Übersetzung"
    assert [c.value for c in ws[1]][:2] == ["Original", "Übersetzung (DE)"]
    assert wb.sheetnames[-1] == "Prüfbericht"
    assert [c.value for c in wb["Prüfbericht"][1]] == ["Schweregrad", "Segment", "Typ", "Meldung"]


def test_the_report_writes_the_findings_in_the_language_of_the_translation() -> None:
    finding = QualityWarning(
        segment_id=None,
        idx=-1,
        lang="en",
        type="missing_glossary",
        code=QualityCode.MISSING_GLOSSARY,
        params={"term": "korpus", "count": "2"},
        severity=Severity.INFO,
        message="Fachbegriff „korpus“ kommt 2× vor, fehlt aber im Glossar.",
    )
    assert report_lines([finding], "en") == [
        "[Hint] Segment -: Technical term “korpus” occurs 2× but is missing from the glossary."
    ]
    assert report_lines([finding], "de") == [
        "[Hinweis] Segment -: Fachbegriff „korpus“ kommt 2× vor, fehlt aber im Glossar."
    ]
    old = finding.model_copy(update={"code": None, "message": "Alte Meldung"})
    assert report_lines([old], "en") == ["[Hint] Segment -: Alte Meldung"]  # stored before the codes existed
    broken = finding.model_copy(update={"params": {}})
    assert report_lines([broken], "en")[0].endswith(finding.message)  # a missing value: the stored message
