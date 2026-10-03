import io
from pathlib import Path

import pymupdf
import pytest
from conftest import Headers, docx_bytes, upload
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen import canvas

from app.config import settings
from app.db import SessionLocal
from app.models import Project
from app.services import parsing

PDF_LINES = ["Assembly guide base unit", "The carcass is made of coated chipboard."]


def pdf_bytes(*lines: str) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i, line in enumerate(lines):
        c.drawString(60, 780 - 28 * i, line)
    c.save()
    return buf.getvalue()


def xlsx_bytes(cells: dict[str, str | float]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for ref, value in cells.items():
        ws[ref] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write(tmp_path: Path, name: str, data: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


# Parsing, unit level
def test_docx_parse_finds_paragraphs_and_table_cells(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Hallo Welt")
    doc.add_paragraph("12345")  # no letters: not translatable
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Zelle eins"
    table.rows[0].cells[1].text = "Zelle zwei"
    path = tmp_path / "t.docx"
    doc.save(str(path))
    parts, structure = parsing.parse_docx(path)
    assert [p.text for p in parts] == ["Hallo Welt", "Zelle eins", "Zelle zwei"]
    assert structure == {"type": "docx", "paragraphs": 3}


def test_docx_write_back_keeps_the_formatting_of_the_first_run(tmp_path: Path) -> None:
    src = write(tmp_path, "src.docx", docx_bytes("Fett gedruckt", bold_first=True))
    parts, _ = parsing.parse_docx(src)
    dst = tmp_path / "dst.docx"
    parsing.write_docx(src, dst, {parts[0].locator["p"]: "Printed in bold"})
    result = Document(str(dst))
    assert [p.text for p in result.paragraphs] == ["Printed in bold"]
    assert result.paragraphs[0].runs[0].bold is True


def test_xlsx_parse_skips_numbers_and_formulas_and_write_back_keeps_them(tmp_path: Path) -> None:
    src = write(tmp_path, "src.xlsx", xlsx_bytes({"A1": "Der Griff", "B1": 5, "C1": "=B1*2", "A2": "12.5 %"}))
    parts, structure = parsing.parse_xlsx(src)
    assert [(p.locator["cell"], p.text) for p in parts] == [("A1", "Der Griff")]
    assert structure["sheets"] == [{"name": "Sheet", "cells": 1}]
    dst = tmp_path / "dst.xlsx"
    parsing.write_xlsx(src, dst, {("Sheet", "A1"): "The handle"})
    ws = load_workbook(dst).active
    assert ws is not None
    assert (ws["A1"].value, ws["B1"].value, ws["C1"].value) == ("The handle", 5, "=B1*2")


def test_pdf_parse_returns_text_blocks_with_page_numbers(tmp_path: Path) -> None:
    parts, structure = parsing.parse_pdf(write(tmp_path, "t.pdf", pdf_bytes(*PDF_LINES)))
    assert [p.text for p in parts] == PDF_LINES
    assert parts[0].locator["page"] == 1
    assert structure == {"type": "pdf", "pages": 1, "blocks": 2}


def test_pdf_without_text_is_reported_as_scanned(tmp_path: Path) -> None:
    empty = io.BytesIO()
    canvas.Canvas(empty).save()
    with pytest.raises(ValueError, match="OCR"):
        parsing.parse_pdf(write(tmp_path, "scan.pdf", empty.getvalue()))


@pytest.mark.parametrize(("name", "expected"), [("a.DOCX", "docx"), ("b.xlsx", "xlsx"), ("c.pdf", "pdf")])
def test_file_type_is_detected_case_insensitively(name: str, expected: str) -> None:
    assert parsing.file_type(name) == expected


@pytest.mark.parametrize(
    ("name", "message"),
    [("a.indd", "InDesign"), ("a.imdd", "InDesign"), ("a.txt", ".txt"), ("noext", "nicht unterstützt")],
)
def test_unsupported_file_types_are_explained(name: str, message: str) -> None:
    with pytest.raises(ValueError, match=message):
        parsing.file_type(name)


def test_signature_check_catches_disguised_files(tmp_path: Path) -> None:
    parsing.check_signature(write(tmp_path, "ok.docx", docx_bytes("x")), "docx")
    parsing.check_signature(write(tmp_path, "ok.pdf", pdf_bytes("x")), "pdf")
    with pytest.raises(ValueError, match="keine gültige DOCX"):
        parsing.check_signature(write(tmp_path, "fake.docx", b"MZ\x90\x00 an executable"), "docx")
    with pytest.raises(ValueError, match="keine gültige PDF"):
        parsing.check_signature(write(tmp_path, "empty.pdf", b""), "pdf")


# Upload through the API
def post_file(client: TestClient, headers: Headers, name: str, data: bytes, **form: str) -> tuple[int, str]:
    r = client.post("/api/projects", headers=headers, files={"file": (name, data)}, data=form)
    return r.status_code, str(r.json().get("detail", ""))


def test_pdf_upload_detects_language_and_defaults_the_target(client: TestClient, translator: Headers) -> None:
    project = upload(
        client, translator, "guide.pdf", pdf_bytes("The carcass is made of the coated chipboard and the oak.")
    )
    assert (project["file_type"], project["detected_lang"], project["target_langs"]) == ("pdf", "en", ["de"])


def test_xlsx_upload_creates_cell_segments(client: TestClient, translator: Headers) -> None:
    project = upload(client, translator, "preise.xlsx", xlsx_bytes({"A1": "Der Griff ist weiß", "A2": 3}))
    detail = client.get(f"/api/projects/{project['id']}", headers=translator).json()
    assert [s["source_text"] for s in detail["segments"]] == ["Der Griff ist weiß"]


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        pytest.param("fake.docx", b"this is not a zip", "keine gültige DOCX", id="not-a-zip"),
        pytest.param("empty.pdf", b"", "keine gültige PDF", id="empty-pdf"),
        ("broken.docx", b"PK\x03\x04 but then garbage", "nicht gelesen"),
        pytest.param("noletters.docx", docx_bytes("12345", "6789"), "kein übersetzbarer Text", id="no-letters"),
        pytest.param(
            "onlyformulas.xlsx", xlsx_bytes({"A1": "=1+1", "B1": 5}), "kein übersetzbarer Text", id="only-formulas"
        ),
        pytest.param("design.indd", b"x", "InDesign", id="indesign"),
        pytest.param("notes.txt", b"hello", "nicht unterstützt", id="text-file"),
    ],
)
def test_bad_uploads_are_rejected_with_a_helpful_message(
    client: TestClient, translator: Headers, name: str, data: bytes, message: str
) -> None:
    status, detail = post_file(client, translator, name, data)
    assert status == 400
    assert message in detail


def test_oversized_upload_is_refused_with_413_and_leaves_no_file_behind(
    client: TestClient, translator: Headers, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "max_upload_mb", 0)
    before = set((settings.data_dir / "uploads").iterdir())
    status, detail = post_file(client, translator, "big.docx", docx_bytes("Hallo Welt"))
    assert status == 413
    assert "zu groß" in detail
    assert set((settings.data_dir / "uploads").iterdir()) == before


def test_invalid_source_language_is_rejected(client: TestClient, translator: Headers) -> None:
    status, _ = post_file(client, translator, "a.docx", docx_bytes("Hallo Welt"), source_lang="klingon")
    assert status == 400


def test_filename_is_reduced_to_its_basename_and_names_are_capped(client: TestClient, translator: Headers) -> None:
    project = upload(
        client,
        translator,
        "../../etc/evil.docx",
        docx_bytes("Hallo Welt, wie geht es dir?"),
        name="x" * 500,
        customer="c" * 500,
    )
    assert project["filename"] == "evil.docx"
    assert len(project["name"]) == 200
    assert len(project["customer"]) == 200
    with SessionLocal() as db:
        stored = Path(db.get(Project, project["id"]).file_path)  # pyright: ignore[reportOptionalMemberAccess]
    assert stored.parent == settings.data_dir / "uploads"  # inside the upload folder, whatever the client sent
    assert "evil" not in stored.name  # random name, the client's filename never reaches the file system


def test_unauthenticated_upload_is_rejected(client: TestClient) -> None:
    assert client.post("/api/projects", files={"file": ("a.docx", docx_bytes("Hallo"))}).status_code == 401


def test_pdf_blocks_survive_a_roundtrip_through_pymupdf(tmp_path: Path) -> None:
    """Guards the local type stub: the real library must keep the API the stub describes."""
    with pymupdf.open(write(tmp_path, "t.pdf", pdf_bytes("Hallo Welt"))) as pdf:
        assert len(pdf) == 1
        x0, y0, x1, y1, text, _, kind = next(iter(pdf)).get_text("blocks")[0]
        assert x0 < x1
        assert y0 < y1
        assert "Hallo Welt" in text
        assert kind == 0
