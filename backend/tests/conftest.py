import io
import os
import tempfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

_tmp = Path(tempfile.mkdtemp())
os.environ.update(
    ENVIRONMENT="test",
    DATABASE_URL=f"sqlite:///{_tmp / 't.db'}",
    DATA_DIR=str(_tmp),
    LLM_PROVIDER="mock",
    LLM_MIN_INTERVAL="0",
    GEMINI_API_KEY="",  # never use a developer's real key (they live in .env, which these variables override)
    GEMINI_MODEL="test-model",
    SECRET_KEY="test-only-secret-key-0123456789-abcdefghij",  # >= 32 bytes, like production requires
    SEED_DEMO_DATA="true",
    LOG_FORMAT="text",
)

import httpx  # noqa: E402
import pytest  # noqa: E402
from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.routers.admin import login_limiter  # noqa: E402

Headers = dict[str, str]


@pytest.fixture(scope="session")
def client() -> Iterator[TestClient]:
    with TestClient(app) as c:
        yield c


@pytest.fixture(autouse=True)
def _no_real_http(monkeypatch: pytest.MonkeyPatch) -> None:
    """Tests must never reach the internet (no quota use, no flakiness). Tests that need HTTP patch this themselves."""

    def refuse(url: str, **_: object) -> None:
        raise AssertionError(f"real HTTP request attempted in a test: {url}")

    monkeypatch.setattr(httpx, "post", refuse)


@pytest.fixture(autouse=True)
def _fresh_login_limiter() -> None:
    login_limiter.clear()  # failures from one test must not lock out the next


def login(client: TestClient, user: str, password: str | None = None) -> Headers:
    r = client.post("/api/auth/login", data={"username": user, "password": password or f"{user}123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="session")
def admin(client: TestClient) -> Headers:
    return login(client, "admin")


@pytest.fixture(scope="session")
def reviewer(client: TestClient) -> Headers:
    return login(client, "reviewer")


@pytest.fixture(scope="session")
def translator(client: TestClient) -> Headers:
    return login(client, "translator")


def docx_bytes(*paragraphs: str, bold_first: bool = False) -> bytes:
    doc = Document()
    for i, text in enumerate(paragraphs):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold_first and i == 0
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload(client: TestClient, headers: Headers, filename: str, data: bytes, **form: str) -> dict[str, Any]:
    r = client.post("/api/projects", headers=headers, files={"file": (filename, data)}, data=form)
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture
def translated_project(client: TestClient, translator: Headers) -> int:
    """A German project that was already translated (preview) to English with the mock provider."""
    text = "Der Korpus und die Frontblende sind weiß. Preis 49,90 EUR, Artikel UK-4711."
    pid: int = upload(client, translator, "fixture.docx", docx_bytes("Montageanleitung", text), target_langs="en")["id"]
    assert client.post(f"/api/projects/{pid}/translate", headers=translator).status_code == 200
    return pid
