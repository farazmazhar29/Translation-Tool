"""Generate fictional demo files in ../samples: `python make_samples.py`."""

from pathlib import Path

from docx import Document
from docx.shared import Cm
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parents[1] / "samples"
OUT.mkdir(exist_ok=True)


def manual() -> None:
    doc = Document()
    doc.add_heading("Montageanleitung Unterschrank", level=1)
    doc.add_paragraph(
        "Die Rotpunkt Kitchen Line überzeugt durch hochwertige Materialien. "
        "Der Korpus besteht aus beschichteter Spanplatte, die Frontblende aus Eiche furniert."
    )
    doc.add_heading("Lieferumfang", level=2)
    for line in (
        "1 Unterschrank mit Auszug, Artikelnummer UK-4711/60",
        "2 Scharniere (Topfband) für die Tür",
        "1 Sockelblende, weiß, Art.-Nr. SB-2200",
        "Bitte den Einlegeboden vor der Montage prüfen.",
    ):
        doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("Montage", level=2)
    doc.add_paragraph(
        "Den Unterschrank ausrichten und mit der Arbeitsplatte verschrauben. Das Gewicht beträgt 34,5 kg."
    )
    doc.add_paragraph("Die Dunstabzugshaube wird erst nach der Nischenrückwand montiert.")
    table = doc.add_table(rows=3, cols=3, style="Light Grid Accent 1")
    for row, values in zip(
        table.rows,
        [
            ("Artikel", "Bezeichnung", "Preis"),
            ("UK-4711/60", "Unterschrank 60 cm", "249,00 EUR"),
            ("HS-9001", "Hängeschrank mit Glastür", "189,90 EUR"),
        ],
        strict=True,
    ):
        for cell, v in zip(row.cells, values, strict=True):
            cell.text = v
    doc.save(OUT / "Montageanleitung_Unterschrank.docx")


def prices() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Preisliste"
    ws.append(["Artikel", "Bezeichnung", "Farbe", "Preis"])
    ws.append(["UK-4711/60", "Unterschrank mit Auszug", "Eiche", 249.0])
    ws.append(["HS-9001", "Hängeschrank mit Glastür", "Weiß", 189.9])
    ws.append(["AP-3300", "Arbeitsplatte, 3 m", "Eiche", 420.0])
    ws.append(["Summe", None, None, "=SUM(D2:D4)"])
    wb.save(OUT / "Preisliste.xlsx")


def glossary() -> None:
    """Terms of the German PDF sample, plus rows that show each import rule (draft, fallbacks, skipped rows)."""
    rows: list[list[str]] = [
        ["de", "en", "de_synonyme", "en_synonyme", "beschreibung", "nicht_uebersetzen", "kategorie", "status"],
        [
            "Arbeitsplatte",
            "worktop",
            "Küchenarbeitsplatte",
            "countertop",
            "Arbeitsfläche über den Unterschränken",
            "",
            "technical",
            "approved",
        ],
        [
            "Frontblende",
            "door front",
            "Front; Blende",
            "front",
            "Sichtbare Vorderseite von Tür oder Schublade",
            "",
            "technical",
            "approved",
        ],
        [
            "beschichtete Spanplatte",
            "melamine-faced chipboard",
            "melaminbeschichtete Spanplatte",
            "",
            "Plattenwerkstoff des Korpus",
            "",
            "technical",
            "approved",
        ],
        ["Dübel", "wall plug", "Wanddübel", "", "Befestigungsmittel für die Wandmontage", "", "technical", "approved"],
        ["Scharnier", "hinge", "Möbelscharnier", "", "", "", "technical", "approved"],
        [
            "Küchenfachberater",
            "kitchen design consultant",
            "Küchenberater",
            "",
            "Ansprechpartner im Fachhandel",
            "",
            "technical",
            "approved",
        ],
        ["Mehrwertsteuer", "VAT", "MwSt.", "value added tax", "", "", "technical", "approved"],
        ["HS-9001", "", "", "", "Artikelnummer Hängeschrank, bleibt unverändert", "ja", "product", "approved"],
        ["Sockelblende", "plinth", "Sockel", "kickboard", "Blende unter den Unterschränken", "", "technical", "draft"],
        [
            "Glastür",
            "glass door",
            "",
            "",
            "Unbekannte Kategorie und Status: technical / approved",
            "",
            "irgendwas",
            "unbekannt",
        ],
        ["Seitenwand", "", "", "", "Übersetzung fehlt: wird übersprungen", "", "technical", "approved"],
        [
            "Arbeitsplatte",
            "work surface",
            "",
            "",
            "Zweites Mal in der Datei: wird übersprungen",
            "",
            "technical",
            "approved",
        ],
    ]
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Glossar"
    for row in rows:
        ws.append([value or None for value in row])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGH", [26, 28, 30, 20, 52, 18, 12, 12], strict=True):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    wb.save(OUT / "glossar_beispiel.xlsx")


def german_pdf() -> None:
    """One page of German assembly instructions (text, not a scan) to try the PDF translation with."""
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=11, leading=16, spaceAfter=10)
    lines: list[tuple[str, str]] = [
        ("Title", "Montageanleitung Hängeschrank"),
        ("Heading2", "Produktbeschreibung"),
        (
            "body",
            "Der Hängeschrank Modell HS-9001 aus der Rotpunkt Kitchen Line überzeugt durch hochwertige Materialien. "
            "Der Korpus besteht aus beschichteter Spanplatte, die Frontblende aus Eiche furniert.",
        ),
        ("Heading2", "Lieferumfang"),
        ("body", "1 Hängeschrank mit Glastür, Artikelnummer HS-9001/60"),
        ("body", "2 Scharniere (Topfband) und 1 Einlegeboden"),
        ("body", "1 Befestigungsschiene, Art.-Nr. BS-3300"),
        ("Heading2", "Montage"),
        (
            "body",
            "Die Befestigungsschiene waagerecht an der Wand anbringen und den Hängeschrank einhängen. "
            "Das Gewicht des Schranks beträgt 28,5 kg. Bitte nur geeignete Dübel verwenden.",
        ),
        (
            "body",
            "Die Dunstabzugshaube wird erst nach der Nischenrückwand montiert. "
            "Die Arbeitsplatte vor dem Einbau auf Beschädigungen prüfen.",
        ),
        ("Heading2", "Hinweise"),
        ("body", "Vor der Montage den Einlegeboden prüfen. Der Preis beträgt 189,90 EUR inklusive Mehrwertsteuer."),
        ("body", "Bei Fragen wenden Sie sich bitte an Ihren Küchenfachberater."),
    ]
    flow: list[Paragraph | Spacer] = [Paragraph(text, styles[kind] if kind != "body" else body) for kind, text in lines]
    SimpleDocTemplate(str(OUT / "Montageanleitung_Haengeschrank.pdf"), pagesize=A4, title="Montageanleitung").build(
        flow
    )


def care_pdf() -> None:
    """One full page of care instructions: paragraphs under headings, as a PDF with real text."""
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10.5, leading=15, spaceAfter=7)
    heading = ParagraphStyle("heading", parent=styles["Heading2"], spaceBefore=8, spaceAfter=5)
    lines: list[tuple[str, str]] = [
        ("Title", "Pflege- und Reinigungsanleitung Küchenmöbel"),
        (
            "body",
            "Vielen Dank, dass Sie sich für eine Küche der Rotpunkt Kitchen Line entschieden haben. Mit der richtigen "
            "Pflege bleiben Fronten, Arbeitsplatte und Beschläge über viele Jahre schön und funktionsfähig. Bitte "
            "bewahren Sie diese Anleitung zusammen mit Ihren Kaufunterlagen auf.",
        ),
        ("heading", "Allgemeine Hinweise"),
        (
            "body",
            "Verwenden Sie zur Reinigung ein weiches, sauberes Tuch oder einen Schwamm und lauwarmes Wasser mit etwas "
            "mildem Spülmittel. Wischen Sie anschließend mit einem trockenen Tuch nach, damit keine Feuchtigkeit in "
            "die Kanten eindringt.",
        ),
        (
            "body",
            "Scheuermittel, Stahlwolle und Dampfreiniger sowie lösungsmittel- oder säurehaltige Reiniger dürfen nicht "
            "verwendet werden. Sie können die Oberflächen dauerhaft beschädigen.",
        ),
        ("heading", "Fronten und Korpus"),
        (
            "body",
            "Die Frontblende besteht je nach Modell aus Lack, Folie oder Echtholzfurnier. Lackierte Fronten reinigen "
            "Sie ohne Druck in Längsrichtung. Furnierte Fronten aus Eiche sollten zweimal im Jahr mit einem geeigneten "
            "Möbelpflegeöl behandelt werden.",
        ),
        (
            "body",
            "Der Korpus aus beschichteter Spanplatte ist feuchtigkeitsbeständig, aber nicht wasserfest. Entfernen Sie "
            "Wasser an Schnittkanten und rund um die Spüle bitte sofort.",
        ),
        (
            "body",
            "Öffnen Sie Geschirrspüler und Backofen nach dem Betrieb erst, wenn der Dampf abgezogen ist. Heißer Dampf "
            "greift die angrenzenden Fronten, die Arbeitsplatte und die Sockelblende an.",
        ),
        ("heading", "Arbeitsplatte"),
        (
            "body",
            "Schneiden Sie nie direkt auf der Arbeitsplatte, sondern verwenden Sie immer ein Schneidebrett. Heiße "
            "Töpfe und Pfannen stellen Sie bitte auf einen Untersetzer.",
        ),
        (
            "body",
            "Arbeitsplatten aus Massivholz werden vor der ersten Benutzung und danach regelmäßig geölt. Für "
            "Arbeitsplatten aus Naturstein empfehlen wir eine jährliche Imprägnierung.",
        ),
        ("heading", "Glastüren und Beschläge"),
        (
            "body",
            "Die Glastür des Hängeschranks HS-9001 reinigen Sie mit handelsüblichem Glasreiniger. Sprühen Sie den "
            "Reiniger auf das Tuch und nicht direkt auf das Glas, damit keine Flüssigkeit in den Rahmen läuft.",
        ),
        (
            "body",
            "Scharniere und Auszüge sind wartungsfrei und dürfen nicht geölt werden. Hängt eine Tür schief oder "
            "schließt sie nicht mehr bündig, lässt sich das Scharnier mit einem Kreuzschlitzschraubendreher nachstellen.",
        ),
        (
            "body",
            "Prüfen Sie einmal im Jahr, ob die Hängeschränke fest an der Wand sitzen. Lockere Dübel oder Schrauben "
            "müssen sofort von einer Fachkraft ersetzt werden.",
        ),
        ("heading", "Service"),
        (
            "body",
            "Bei Fragen zur Pflege oder bei Schäden wenden Sie sich bitte an Ihren Küchenfachberater. Ersatzteile wie "
            "Griffe, Scharniere oder die Sockelblende SB-2200 erhalten Sie über unseren Kundendienst. Alle Preise "
            "enthalten die gesetzliche Mehrwertsteuer.",
        ),
    ]
    own = {"body": body, "heading": heading}
    flow: list[Paragraph | Spacer] = [
        Paragraph(text, own[kind] if kind in own else styles[kind]) for kind, text in lines
    ]
    SimpleDocTemplate(str(OUT / "Pflegeanleitung_Kueche.pdf"), pagesize=A4, title="Pflegeanleitung").build(flow)


def warranty_docx() -> None:
    """One full page of warranty terms: headings, paragraphs, a table and a list, as a Word document."""
    doc = Document()
    page = doc.sections[0]  # python-docx starts with US Letter: A4 with DIN 5008 margins instead
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin, page.bottom_margin, page.left_margin, page.right_margin = Cm(2), Cm(2), Cm(2.5), Cm(2)
    doc.add_heading("Garantie- und Servicebedingungen", level=1)
    doc.add_paragraph(
        "Für alle Möbel der Rotpunkt Kitchen Line gewähren wir zusätzlich zur gesetzlichen Gewährleistung eine "
        "freiwillige Herstellergarantie. Diese Bedingungen gelten für Küchen, die ab dem 1. Januar 2026 gekauft wurden."
    )
    doc.add_heading("1. Umfang der Garantie", level=2)
    doc.add_paragraph(
        "Die Garantie umfasst Material- und Verarbeitungsfehler, die bei bestimmungsgemäßem Gebrauch im privaten "
        "Haushalt auftreten. Wir entscheiden, ob das betroffene Teil repariert oder ersetzt wird. Ausgetauschte Teile "
        "gehen in unser Eigentum über."
    )
    periods = [
        ("Bauteil", "Garantiezeit", "Hinweis"),
        ("Korpus aus beschichteter Spanplatte", "10 Jahre", "nur bei Montage durch einen Fachbetrieb"),
        ("Frontblende und Glastür", "5 Jahre", "ohne natürliche Farbabweichungen von Holz"),
        ("Scharniere und Auszüge", "10 Jahre", "bis zu 80.000 Öffnungen"),
        ("Arbeitsplatte", "2 Jahre", "nicht bei Schnitt- oder Hitzeschäden"),
        ("Elektrogeräte", "2 Jahre", "nach den Bedingungen des Geräteherstellers"),
    ]
    table = doc.add_table(rows=len(periods), cols=3, style="Light Grid Accent 1")
    for row, values in zip(table.rows, periods, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    doc.add_heading("2. Ausschlüsse", level=2)
    doc.add_paragraph("Von der Garantie ausgenommen sind:")
    for line in (
        "Schäden durch unsachgemäße Montage, zum Beispiel durch ungeeignete Dübel oder eine nicht tragfähige Wand",
        "Schäden durch Feuchtigkeit, Hitze oder falsche Reinigungsmittel",
        "gewöhnliche Abnutzung sowie Kratzer und Dellen durch den täglichen Gebrauch",
        "Veränderungen oder Reparaturen, die nicht von uns oder einem Vertragspartner durchgeführt wurden",
    ):
        doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("3. Garantiefall melden", level=2)
    doc.add_paragraph(
        "Melden Sie einen Mangel bitte innerhalb von zwei Monaten, nachdem Sie ihn entdeckt haben. Wenden Sie sich "
        "dazu an Ihren Küchenfachberater oder an unseren Kundendienst und halten Sie die Rechnung, die Auftragsnummer "
        "und ein Foto des Schadens bereit."
    )
    doc.add_paragraph(
        "Bei einem anerkannten Garantiefall tragen wir die Kosten für Ersatzteile, Arbeitszeit und Anfahrt. Liegt "
        "kein Garantiefall vor, erhalten Sie vorab einen Kostenvoranschlag. Alle Preise verstehen sich inklusive "
        "Mehrwertsteuer."
    )
    doc.add_heading("4. Service nach der Garantiezeit", level=2)
    doc.add_paragraph(
        "Auch nach Ablauf der Garantie liefern wir Ersatzteile wie Scharniere, Griffe oder die Sockelblende SB-2200 "
        "mindestens zehn Jahre nach dem Produktionsende des jeweiligen Modells. Den Hängeschrank HS-9001 und den "
        "Unterschrank UK-4711/60 können Sie zudem jederzeit nachbestellen."
    )
    doc.add_paragraph("Kundendienst · Stand: Oktober 2026")
    doc.save(OUT / "Garantiebedingungen.docx")


Part = tuple[str, str, str, str, int, float]  # article number, name, material, fits, packing unit, price

SPARE_PARTS: dict[str, list[Part]] = {
    "Scharniere und Beschläge": [
        ("SC-110", "Topfband-Scharnier 110° mit Dämpfung", "Stahl vernickelt", "alle Drehtüren", 2, 12.90),
        ("SC-170", "Weitwinkel-Scharnier 170°", "Stahl vernickelt", "Eckschränke", 2, 18.50),
        ("MP-001", "Montageplatte für Scharnier", "Stahl", "alle Drehtüren", 2, 3.20),
        ("TD-010", "Türdämpfer zum Nachrüsten", "Kunststoff, grau", "alle Drehtüren", 2, 9.90),
        ("KL-200", "Klappenbeschlag mit Gasdruckfeder", "Stahl", "Hängeschränke mit Klappe", 1, 34.00),
        ("AZ-500", "Vollauszug 500 mm mit Selbsteinzug", "Stahl, grau", "Unterschrank UK-4711/60", 1, 46.00),
        ("AZ-450", "Vollauszug 450 mm mit Selbsteinzug", "Stahl, grau", "Unterschrank 45 cm", 1, 42.00),
        ("SL-001", "Schubkastenschloss mit 2 Schlüsseln", "Zink-Druckguss", "Schubkästen", 1, 12.50),
        ("GR-128", "Bügelgriff 128 mm", "Edelstahl gebürstet", "alle Fronten", 1, 7.90),
        ("GR-160", "Bügelgriff 160 mm", "schwarz matt", "alle Fronten", 1, 8.90),
    ],
    "Fronten, Türen und Blenden": [
        ("FB-6001", "Frontblende 60 × 72 cm", "Eiche furniert", "Unterschrank UK-4711/60", 1, 129.00),
        ("FB-6002", "Frontblende 60 × 72 cm", "Lack weiß matt", "Unterschrank UK-4711/60", 1, 98.00),
        ("GT-9001", "Glastür 60 × 72 cm, Klarglas", "Aluminiumrahmen", "Hängeschrank HS-9001", 1, 149.00),
        ("BB-150", "Blindblende 15 × 72 cm", "Lack weiß matt", "Eckschränke", 1, 32.00),
        ("PL-050", "Passleiste 5 × 72 cm", "Lack weiß matt", "Wandanschluss", 1, 18.00),
        ("SB-2200", "Sockelblende 300 cm", "weiß", "Unterschränke", 1, 39.00),
        ("SB-2201", "Sockelblende 300 cm", "Eiche Dekor", "Unterschränke", 1, 45.00),
        ("KB-300", "Kranzblende 300 cm", "weiß", "Oberseite der Hängeschränke", 1, 42.00),
        ("LB-300", "Lichtblende 300 cm", "weiß", "Unterseite der Hängeschränke", 1, 38.00),
    ],
    "Arbeitsplatte und Nische": [
        ("AP-3300", "Arbeitsplatte 300 × 60 cm", "Eiche massiv, geölt", "Küchenzeilen", 1, 420.00),
        ("AP-3301", "Arbeitsplatte 300 × 60 cm", "Schichtstoff, Betonoptik", "Küchenzeilen", 1, 189.00),
        ("EV-001", "Eckverbindungsbeschlag", "Stahl", "Arbeitsplatten über Eck", 3, 7.50),
        ("KD-001", "Kantendichtung", "Silikon, transparent", "Arbeitsplatte", 1, 11.50),
        ("AL-030", "Abschlussleiste", "Aluminium", "Wandanschluss der Arbeitsplatte", 1, 16.90),
        ("NR-300", "Nischenrückwand 300 × 60 cm", "Glas, weiß", "zwischen Arbeitsplatte und Hängeschrank", 1, 289.00),
        ("SM-080", "Schutzmatte, rutschfest", "Aluminium", "Spülenunterschrank 80 cm", 1, 24.90),
    ],
    "Beleuchtung": [
        ("LE-060", "LED-Unterbauleuchte 60 cm, warmweiß", "Aluminium", "Hängeschränke", 1, 39.90),
        ("LN-030", "LED-Netzteil 30 W", "Kunststoff", "LED-Beleuchtung", 1, 29.00),
        ("FS-001", "Funkschalter mit Dimmfunktion", "weiß", "LED-Beleuchtung", 1, 19.90),
    ],
    "Montagematerial": [
        ("BS-3300", "Befestigungsschiene 100 cm", "Stahl verzinkt", "Hängeschränke", 1, 19.90),
        ("AH-100", "Schrankaufhänger, verstellbar", "Stahl und Kunststoff", "Hängeschränke", 2, 8.60),
        ("DB-080", "Dübel 8 mm mit Schraube", "Nylon und Stahl", "Beton und Vollziegel", 10, 4.90),
        ("DB-081", "Hohlraumdübel 8 mm", "Metall", "Gipskartonwände", 10, 9.80),
        ("SF-100", "Stellfuß, höhenverstellbar 100–150 mm", "Kunststoff", "Unterschränke", 4, 6.80),
        ("SK-010", "Sockelklammer für Stellfuß", "Kunststoff", "Sockelblende", 4, 3.60),
        ("EB-560", "Einlegeboden 56 × 30 cm", "beschichtete Spanplatte, weiß", "Hängeschrank HS-9001", 1, 14.50),
        ("EB-561", "Einlegeboden 56 × 30 cm", "Glas, satiniert", "Hängeschrank HS-9001", 1, 24.90),
        ("BT-200", "Bodenträger mit Kippsicherung", "Kunststoff, transparent", "alle Einlegeböden", 4, 2.40),
        ("VS-050", "Verbindungsschraube", "Stahl verzinkt", "Korpusse nebeneinander", 8, 3.90),
        ("AK-100", "Abdeckkappen für Schrauben", "Kunststoff, weiß", "sichtbare Schrauben", 20, 2.50),
    ],
    "Pflege und Reparatur": [
        ("PF-250", "Möbelpflegeöl 250 ml", "farblos", "Holzfronten und Massivholz", 1, 14.90),
        ("RG-500", "Glasreiniger 500 ml", "farblos", "Glastüren und Nischenrückwand", 1, 6.90),
        ("MT-002", "Mikrofasertuch, 2 Stück", "grau", "alle Oberflächen", 1, 8.50),
        ("RS-001", "Reparaturset für Kratzer", "Wachsstifte, 6 Farben", "beschichtete Oberflächen", 1, 19.90),
    ],
}


def spare_parts_xlsx() -> None:
    """One printed page of a spare parts list: text cells are translated, article numbers and prices stay."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Ersatzteile"
    ws.append(["Ersatzteilliste Rotpunkt Kitchen Line"])
    ws.append(["Stand: Oktober 2026 · Alle Preise in EUR inklusive Mehrwertsteuer"])
    ws.append([])
    ws.append(["Art.-Nr.", "Bezeichnung", "Material / Farbe", "Passend für", "VE", "Preis"])
    bold = [1, 4]  # title and column heads
    for group, parts in SPARE_PARTS.items():
        ws.append([group])
        bold.append(ws.max_row)
        for part in parts:
            ws.append(list(part))
            ws.cell(ws.max_row, 6).number_format = "#,##0.00"
    ws.append([])
    ws.append(["Hinweise"])
    bold.append(ws.max_row)
    for note in (
        "VE = Verpackungseinheit. Scharniere und Schrankaufhänger werden paarweise geliefert.",
        "Bestellungen nimmt Ihr Küchenfachberater entgegen; die Lieferzeit beträgt etwa zwei Wochen.",
        "Farbabweichungen bei Echtholzfronten sind naturbedingt und kein Reklamationsgrund.",
    ):
        ws.append([note])
    for row in bold:
        for cell in ws[row]:
            cell.font = Font(bold=True, size=14 if row == 1 else 11)
    ws["A2"].font = Font(italic=True)
    for column, width in zip("ABCDEF", [10, 36, 28, 36, 5, 10], strict=True):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(OUT / "Ersatzteilliste.xlsx")


if __name__ == "__main__":
    manual()
    prices()
    glossary()
    german_pdf()
    care_pdf()
    warranty_docx()
    spare_parts_xlsx()
