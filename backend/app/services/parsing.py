"""Split uploaded files into segments (paragraph / cell / text block) and write translations back.

Every segment has a `locator` so the translation lands exactly where the original text was.
"""

import re
from pathlib import Path
from typing import Any, NamedTuple, cast

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from ..enums import FileFormat

FILE_TYPES = {".docx": FileFormat.DOCX, ".xlsx": FileFormat.XLSX, ".pdf": FileFormat.PDF}
LETTER = re.compile(r"[A-Za-zÄÖÜäöüß]")


class Part(NamedTuple):
    locator: dict[str, Any]
    text: str


def file_type(filename: str) -> FileFormat:
    ext = Path(filename).suffix.lower()
    if ext in (".imdd", ".indd"):
        raise ValueError("InDesign-Dateien werden erst in Phase 2 unterstützt.")
    if ext not in FILE_TYPES:
        raise ValueError(f"Dateityp „{ext}“ nicht unterstützt (erlaubt: docx, xlsx, pdf).")
    return FILE_TYPES[ext]


SIGNATURES = {
    FileFormat.DOCX: b"PK\x03\x04",
    FileFormat.XLSX: b"PK\x03\x04",
    FileFormat.PDF: b"%PDF-",
}  # Office files are zip archives


def check_signature(path: Path, ftype: FileFormat) -> None:
    """The extension alone proves nothing: make sure the content looks like what it claims to be."""
    with path.open("rb") as f:
        head = f.read(8)
    if not head.startswith(SIGNATURES[ftype]):
        raise ValueError(f"Die Datei ist keine gültige {ftype.upper()}-Datei (Inhalt passt nicht zur Endung).")


def _translatable(text: str) -> bool:
    return bool(LETTER.search(text))


# docx
def _paragraphs(path: Path) -> tuple[Any, list[Paragraph]]:
    doc = Document(str(path))
    return doc, [Paragraph(cast(CT_P, p), doc) for p in doc.element.iter(qn("w:p"))]  # incl. table cells


def parse_docx(path: Path) -> tuple[list[Part], dict[str, Any]]:
    parts = [
        Part({"type": FileFormat.DOCX, "p": i}, p.text)
        for i, p in enumerate(_paragraphs(path)[1])
        if _translatable(p.text)
    ]
    return parts, {"type": FileFormat.DOCX, "paragraphs": len(parts)}


def write_docx(src: Path, dst: Path, repl: dict[int, str]) -> None:
    doc, paras = _paragraphs(src)
    for i, text in repl.items():
        runs = paras[i].runs
        if runs:  # the first run keeps the original formatting
            runs[0].text = text
            for r in runs[1:]:
                r.text = ""
        else:
            paras[i].add_run(text)
    doc.save(str(dst))


# xlsx
def parse_xlsx(path: Path) -> tuple[list[Part], dict[str, Any]]:
    parts: list[Part] = []
    sheets: list[dict[str, Any]] = []
    for ws in load_workbook(path).worksheets:
        found = [
            Part({"type": FileFormat.XLSX, "sheet": ws.title, "cell": c.coordinate}, c.value)
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and not c.value.startswith("=") and _translatable(c.value)
        ]
        parts += found
        sheets.append({"name": ws.title, "cells": len(found)})
    return parts, {"type": FileFormat.XLSX, "sheets": sheets}


def write_xlsx(src: Path, dst: Path, repl: dict[tuple[str, str], str]) -> None:
    wb = load_workbook(src)
    for (sheet, cell), text in repl.items():
        wb[sheet][cell].value = text
    wb.save(dst)


# pdf
def parse_pdf(path: Path) -> tuple[list[Part], dict[str, Any]]:
    parts: list[Part] = []
    with pymupdf.open(path) as pdf:
        for n, page in enumerate(pdf, start=1):
            for x0, y0, x1, y1, raw, _, kind in page.get_text("blocks"):
                text = re.sub(r"\s*\n\s*", " ", str(raw)).strip()
                if kind == 0 and _translatable(text):
                    parts.append(Part({"type": FileFormat.PDF, "page": n, "bbox": [x0, y0, x1, y1]}, text))
        pages = len(pdf)
    if not parts:
        raise ValueError("Im PDF wurde kein Text erkannt (gescannte PDFs: OCR folgt in Phase 2).")
    return parts, {"type": FileFormat.PDF, "pages": pages, "blocks": len(parts)}


def parse_file(path: Path, ftype: FileFormat) -> tuple[list[Part], dict[str, Any]]:
    return {FileFormat.DOCX: parse_docx, FileFormat.XLSX: parse_xlsx, FileFormat.PDF: parse_pdf}[ftype](path)
