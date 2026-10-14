"""Export approved translations as Word, Excel or PDF.

mode   "translation" = translation only | "both" = original + translation
opts   include_report (quality report), include_comments
config Excel layout ("Erweiterte Einstellungen"): free choice of source/target/comment/position columns
The fixed words (headings, the quality report) are in the language of the translation, see `texts`.
"""

import re
import uuid
from collections.abc import Callable
from pathlib import Path
from typing import Any, NamedTuple
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell  # pyright: ignore[reportPrivateUsage]
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string
from openpyxl.utils.cell import coordinate_from_string
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import PropertySet, getSampleStyleSheet
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.config import settings
from app.core.enums import ExportMode, FileFormat
from app.export.texts import finding_text, labels, severity
from app.projects import parsing
from app.projects.models import Project
from app.quality.schemas import QualityWarning

XLSX_DEFAULTS: dict[str, Any] = {
    "source_col": "A",
    "target_col": "B",
    "comment_col": "C",
    "location_col": "",
    "start_row": 2,
    "header": True,
}
Exporter = Callable[[Project, str, ExportMode, dict[str, bool], list[QualityWarning], dict[str, Any] | None], Path]


def rows(p: Project, lang: str) -> list[dict[str, Any]]:
    words = labels(lang)

    def row(seg_no: int, loc: dict[str, Any], src: str, text: str, comment: str) -> dict[str, Any]:
        place = {
            FileFormat.XLSX: f"{loc.get('sheet')}!{loc.get('cell')}",
            FileFormat.PDF: words["page"].format(n=loc.get("page")),
        }.get(loc["type"], words["paragraph"].format(n=loc.get("p", 0) + 1))
        return {"no": seg_no, "locator": loc, "source": src, "target": text, "comment": comment, "location": place}

    return [
        row(s.idx + 1, s.locator, s.source_text, t.text if t else "", t.comment if t else "")
        for s in p.segments
        for t in [s.translation(lang)]
    ]


def out_path(ext: FileFormat) -> Path:
    """Stored exports get a unique name: two projects with the same filename must never overwrite each other."""
    return settings.data_dir / "exports" / f"{uuid.uuid4().hex}.{ext}"


def download_name(p: Project, lang: str, mode: ExportMode, ext: FileFormat) -> str:
    """What the user sees when downloading (sanitised, since the original filename is user input)."""
    stem = re.sub(r"[^\w.-]+", "_", Path(p.filename).stem).strip("._") or f"projekt-{p.id}"
    return f"{stem[:80]}_{lang}_{mode}.{ext}"


def report_lines(warnings: list[QualityWarning], lang: str) -> list[str]:
    segment = labels(lang)["segment"]
    return [
        f"[{severity(w.severity, lang)}] {segment} {w.idx + 1 if w.idx >= 0 else '-'}: {finding_text(w, lang)}"
        for w in warnings
    ]


# An Excel file exported as Word or PDF stays a table: its rows and columns, the text cells translated.
class SheetCell(NamedTuple):
    text: str
    original: str = ""  # "both": the original below the translation
    comment: str = ""
    bold: bool = False
    number: bool = False


class Sheet(NamedTuple):
    name: str
    rows: list[list[SheetCell | None]]
    spans: set[int]  # rows with a single cell in the first column (a title or a group): it spans the table
    numeric: set[int]  # columns holding numbers: right-aligned, their heads too


def right_aligned(sheet: Sheet, col: int, value: SheetCell) -> bool:
    return value.number or col in sheet.numeric


def column_widths(sheet: Sheet, width: float, size: float) -> list[float]:
    """Widths in points: every column at least as wide as its longest word (a price or an article number must not
    break), the rest of the line shared by how much text a column holds."""
    cols = range(len(sheet.rows[0]))
    cells = [[v for i, line in enumerate(sheet.rows) if i not in sheet.spans and (v := line[c])] for c in cols]
    pad = 2 * size  # cell padding left and right

    def word(v: SheetCell) -> float:
        font = "Helvetica-Bold" if v.bold else "Helvetica"
        return max((stringWidth(w, font, size) for w in v.text.split()), default=0)

    least = [max([3 * size, *(word(v) for v in col)]) + pad for col in cells]
    share = [max([4, *(min(40, len(v.text)) for v in col)]) for col in cells]
    rest = max(0.0, width - sum(least))
    return [low + rest * s / sum(share) for low, s in zip(least, share, strict=True)]


def number_text(value: float, decimals: int, lang: str) -> str:
    """12.9 with 2 decimals: "12,90" in German, "12.90" in English, as the reader of the translation writes it."""
    text = f"{value:,.{decimals}f}"
    return text.translate(str.maketrans(",.", ".,")) if lang == "de" else text


def sheets(p: Project, lang: str, both: bool, comments: bool) -> list[Sheet]:
    """The sheets of the original file, every text cell replaced by its translation; empty rows are left out."""
    by_cell = {(str(s.locator.get("sheet")), str(s.locator.get("cell"))): s for s in p.segments}
    out: list[Sheet] = []
    for sheet in parsing.read_cells(p.original()):
        placed: dict[tuple[int, int], SheetCell] = {}
        for c in sheet["cells"]:
            letters, row = coordinate_from_string(c["ref"])
            if c["number"] is None:
                seg = by_cell.get((sheet["name"], c["ref"]))
                tr = seg.translation(lang) if seg else None
                text = tr.text if tr and tr.text else c["text"]
                original = seg.source_text if both and seg and seg.source_text != text else ""
                comment = tr.comment if comments and tr else ""
                cell = SheetCell(text, original, comment, c["bold"])
            else:
                cell = SheetCell(number_text(c["number"], c["decimals"], lang), bold=c["bold"], number=True)
            placed[(row, column_index_from_string(letters))] = cell
        if not placed:
            continue
        width = max(col for _, col in placed)
        lines = [[placed.get((row, col)) for col in range(1, width + 1)] for row in sorted({row for row, _ in placed})]
        spans = {i for i, line in enumerate(lines) if line[0] and sum(c is not None for c in line) == 1}
        numeric = {col - 1 for (_, col), cell in placed.items() if cell.number}
        out.append(Sheet(sheet["name"], lines, spans, numeric))
    return out


# Word
GREY = RGBColor(0x6B, 0x6B, 0x6B)


def _docx_cell(cell: _Cell, value: SheetCell, comment_label: str, right: bool = False) -> None:
    first = cell.paragraphs[0]
    run = first.add_run(value.text)
    run.bold, run.font.size = value.bold, Pt(9)
    if value.original:
        small = cell.add_paragraph().add_run(value.original)
        small.font.size, small.font.color.rgb = Pt(7.5), GREY
    if value.comment:
        note = cell.add_paragraph().add_run(f"{comment_label}: {value.comment}")
        note.italic, note.font.size = True, Pt(7.5)
    if right:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _docx_sheets(doc: Any, tables: list[Sheet], comment_label: str) -> None:
    # A4 instead of python-docx's US Letter; a sheet with many columns gets a landscape page
    page = doc.sections[0]
    wide = any(len(s.rows[0]) > 6 for s in tables)
    page.page_width, page.page_height = (Cm(29.7), Cm(21)) if wide else (Cm(21), Cm(29.7))
    if wide:
        page.orientation = WD_ORIENT.LANDSCAPE
    page.left_margin = page.right_margin = page.top_margin = page.bottom_margin = Cm(2)
    usable = (page.page_width - page.left_margin - page.right_margin) / 12700  # EMU -> points
    for sheet in tables:
        if len(tables) > 1:
            doc.add_heading(sheet.name, level=2)
        table = doc.add_table(rows=0, cols=len(sheet.rows[0]))
        table.style, table.autofit = "Table Grid", False
        widths = [Pt(w) for w in column_widths(sheet, usable, 9)]
        for i, line in enumerate(sheet.rows):
            cells = table.add_row().cells
            for cell, w in zip(cells, widths, strict=True):
                cell.width = w  # Word follows the cell widths, not the column's
            if i in sheet.spans and line[0]:
                _docx_cell(cells[0].merge(cells[-1]), line[0], comment_label)
                continue
            for col, (cell, value) in enumerate(zip(cells, line, strict=True)):
                if value:
                    _docx_cell(cell, value, comment_label, right_aligned(sheet, col, value))


def to_docx(
    p: Project,
    lang: str,
    mode: ExportMode,
    opts: dict[str, bool],
    warnings: list[QualityWarning],
    _config: dict[str, Any] | None = None,
) -> Path:
    data, path = rows(p, lang), out_path(FileFormat.DOCX)
    if mode == ExportMode.TRANSLATION and p.file_type == FileFormat.DOCX and not any(opts.values()):
        parsing.write_docx(p.original(), path, {r["locator"]["p"]: r["target"] for r in data if r["target"]})
        return path  # original layout and formatting stay intact
    words = labels(lang)
    doc = Document()
    doc.add_heading(f"{p.name} ({lang.upper()})", level=1)
    if p.file_type == FileFormat.XLSX:  # a sheet stays a table
        _docx_sheets(doc, sheets(p, lang, mode == ExportMode.BOTH, opts["include_comments"]), words["comment"])
    elif mode == ExportMode.BOTH:
        head = ["#", words["original"], words["translation"]] + ([words["comment"]] if opts["include_comments"] else [])
        table = doc.add_table(rows=1, cols=len(head), style="Light Grid Accent 1")
        for cell, h in zip(table.rows[0].cells, head, strict=False):
            cell.text = h
        for r in data:
            values = [str(r["no"]), r["source"], r["target"], r["comment"]]
            for cell, v in zip(table.add_row().cells, values, strict=False):
                cell.text = v
    else:
        for r in data:
            doc.add_paragraph(r["target"])
            if opts["include_comments"] and r["comment"]:
                run = doc.add_paragraph().add_run(f"{words['comment']}: {r['comment']}")
                run.italic, run.font.size = True, Pt(9)
    if opts["include_report"]:
        doc.add_page_break()
        doc.add_heading(words["report"], level=1)
        for line in report_lines(warnings, lang) or [words["no_findings"]]:
            doc.add_paragraph(line, style="List Bullet")
    doc.save(str(path))
    return path


# Excel
def to_xlsx(
    p: Project,
    lang: str,
    mode: ExportMode,
    opts: dict[str, bool],
    warnings: list[QualityWarning],
    config: dict[str, Any] | None = None,
) -> Path:
    data, path = rows(p, lang), out_path(FileFormat.XLSX)
    if mode == ExportMode.TRANSLATION and p.file_type == FileFormat.XLSX and not config and not opts["include_report"]:
        parsing.write_xlsx(
            p.original(),
            path,
            {(r["locator"]["sheet"], r["locator"]["cell"]): r["target"] for r in data if r["target"]},
        )
        return path
    words = labels(lang)
    cfg = {**XLSX_DEFAULTS, "sheet_name": words["sheet"], **(config or {})}
    wb = Workbook()
    ws = wb.active or wb.create_sheet()
    ws.title = str(cfg["sheet_name"])[:31] or words["sheet"]
    layout = [(str(cfg["source_col"]), words["original"], "source")] if mode == ExportMode.BOTH else []
    layout.append((str(cfg["target_col"]), f"{words['translation']} ({lang.upper()})", "target"))
    if opts["include_comments"] and cfg["comment_col"]:
        layout.append((str(cfg["comment_col"]), words["comment"], "comment"))
    if cfg["location_col"]:
        layout.append((str(cfg["location_col"]), words["position"], "location"))
    start = int(cfg["start_row"])
    for col, head, field in layout:
        ws.column_dimensions[col.upper()].width = 60
        if cfg["header"] and start > 1:
            ws.cell(row=start - 1, column=column_index_from_string(col.upper()), value=head).font = Font(bold=True)
        for i, r in enumerate(data):
            ws.cell(row=start + i, column=column_index_from_string(col.upper()), value=r[field])
    if opts["include_report"]:
        rs = wb.create_sheet(words["report"][:31])
        rs.append([words["severity"], words["segment"], words["type"], words["message"]])
        rs.column_dimensions["D"].width = 90
        for w in warnings:
            rs.append([severity(w.severity, lang), w.idx + 1 if w.idx >= 0 else "", w.type, finding_text(w, lang)])
    wb.save(path)
    return path


# PDF
MARGIN = 30


def _pdf_sheets(tables: list[Sheet], width: float, comment_label: str, heading: PropertySet) -> list[Any]:
    """Each sheet as a table, the columns sized by `column_widths`."""
    base = getSampleStyleSheet()["BodyText"]
    left, right = base.clone("cell", fontSize=8.5, leading=10.5), base.clone("num", fontSize=8.5, leading=10.5)
    right.alignment = TA_RIGHT

    def cell(v: SheetCell | None, align_right: bool) -> Any:
        if v is None:
            return ""
        html = f"<b>{escape(v.text)}</b>" if v.bold else escape(v.text)
        if v.original:
            html += f'<br/><font size="7" color="#6b6b6b">{escape(v.original)}</font>'
        if v.comment:
            html += f'<br/><font size="7"><i>{escape(comment_label)}: {escape(v.comment)}</i></font>'
        return Paragraph(html, right if align_right else left)

    story: list[Any] = []
    for sheet in tables:
        if len(tables) > 1:
            story += [Spacer(1, 10), Paragraph(escape(sheet.name), heading)]

        data = [
            [cell(v, i not in sheet.spans and v is not None and right_aligned(sheet, c, v)) for c, v in enumerate(line)]
            for i, line in enumerate(sheet.rows)
        ]
        style = [
            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            *[("SPAN", (0, i), (-1, i)) for i in sheet.spans],
        ]
        table = Table(data, colWidths=column_widths(sheet, width, 8.5))
        table.setStyle(TableStyle(style))
        story.append(table)
    return story


def to_pdf(
    p: Project,
    lang: str,
    mode: ExportMode,
    opts: dict[str, bool],
    warnings: list[QualityWarning],
    _config: dict[str, Any] | None = None,
) -> Path:
    data, path = rows(p, lang), out_path(FileFormat.PDF)
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=9, leading=11)

    def para(text: str, style: PropertySet = small) -> Paragraph:
        return Paragraph(escape(text), style)

    words = labels(lang)
    story: list[Any] = [para(f"{p.name} ({lang.upper()})", styles["Title"])]
    page = A4
    if p.file_type == FileFormat.XLSX:  # a sheet stays a table; a wide one gets a landscape page
        tables = sheets(p, lang, mode == ExportMode.BOTH, opts["include_comments"])
        page = landscape(A4) if any(len(s.rows[0]) > 6 for s in tables) else A4
        story += _pdf_sheets(tables, page[0] - 2 * MARGIN, words["comment"], styles["Heading2"])
    elif mode == ExportMode.BOTH:
        comments = opts["include_comments"]
        table: list[list[Any]] = [
            ["#", words["original"], words["translation"]] + ([words["comment"]] if comments else [])
        ]
        table += [
            [str(r["no"]), para(r["source"]), para(r["target"])] + ([para(r["comment"])] if comments else [])
            for r in data
        ]
        t = Table(table, colWidths=[25, 210, 210] + ([80] if comments else []), repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ]
            )
        )
        story.append(t)
    else:
        for r in data:
            story.append(para(r["target"], styles["BodyText"]))
            if opts["include_comments"] and r["comment"]:
                story.append(para(f"{words['comment']}: {r['comment']}"))
    if opts["include_report"]:
        story += [Spacer(1, 20), para(words["report"], styles["Heading1"])]
        story += [para(line) for line in report_lines(warnings, lang) or [words["no_findings"]]]
    SimpleDocTemplate(
        str(path), pagesize=page, leftMargin=MARGIN, rightMargin=MARGIN, topMargin=MARGIN, bottomMargin=MARGIN
    ).build(story)
    return path


EXPORTERS: dict[FileFormat, Exporter] = {
    FileFormat.DOCX: to_docx,
    FileFormat.XLSX: to_xlsx,
    FileFormat.PDF: to_pdf,
}
