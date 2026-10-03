import io
import re
from typing import Annotated

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook

from ..auth import Admin, Me, Reviewer, Translator
from ..db import DB
from ..models import GlossaryEntry, GlossaryVersion
from ..schemas import DuplicateGroup, GlossaryIn, GlossaryOut, GlossaryVersionOut, ImportResult
from ..services.translation import LANGS

router = APIRouter(prefix="/api/glossary", tags=["Glossary"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COLUMNS = ["de", "en", "de_synonyme", "en_synonyme", "beschreibung", "nicht_uebersetzen", "kategorie", "status"]
ALIASES = {
    "deutsch": "de",
    "german": "de",
    "englisch": "en",
    "english": "en",
    "description": "beschreibung",
    "dnt": "nicht_uebersetzen",
    "do_not_translate": "nicht_uebersetzen",
    "category": "kategorie",
    "de_synonyms": "de_synonyme",
    "en_synonyms": "en_synonyme",
}


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower())


def duplicates(db: DB, terms: dict[str, str], skip: int | None = None) -> list[GlossaryEntry]:
    wanted = {(lang, norm(t)) for lang, t in terms.items()}
    return [
        e
        for e in db.query(GlossaryEntry)
        if e.id != skip
        and wanted
        & (
            {(lang, norm(t)) for lang, t in e.terms.items()}
            | {(lang, norm(s)) for lang, ss in e.synonyms.items() for s in ss}
        )
    ]


def snapshot(db: DB, e: GlossaryEntry, action: str, user: str) -> None:
    db.add(
        GlossaryVersion(
            entry_id=e.id,
            entry_version=e.version,
            action=action,
            changed_by=user,
            data={
                "terms": e.terms,
                "synonyms": e.synonyms,
                "description": e.description,
                "category": e.category,
                "do_not_translate": e.do_not_translate,
                "status": e.status,
            },
        )
    )


def clean(body: GlossaryIn) -> tuple[dict[str, str], dict[str, list[str]]]:
    terms = {k: v.strip() for k, v in body.terms.items() if k in LANGS and v.strip()}
    if not terms or (not body.do_not_translate and len(terms) < 2):
        raise HTTPException(400, "Begriff und Übersetzung angeben (oder „Nicht übersetzen“ aktivieren).")
    return terms, {k: [s.strip() for s in v if s.strip()] for k, v in body.synonyms.items() if k in LANGS}


def conflict(dups: list[GlossaryEntry]) -> HTTPException:
    found = [GlossaryOut.model_validate(d).model_dump(mode="json") for d in dups]
    return HTTPException(409, {"message": "Dublette gefunden", "duplicates": found})


def find(db: DB, eid: int) -> GlossaryEntry:
    e = db.get(GlossaryEntry, eid)
    if e is None:
        raise HTTPException(404, "Eintrag nicht gefunden")
    return e


@router.get("", response_model=list[GlossaryOut])
def list_entries(db: DB, _: Me, q: str = "", status: str = "", category: str = "") -> list[GlossaryEntry]:
    def text(e: GlossaryEntry) -> str:
        return " ".join([*e.terms.values(), *(s for v in e.synonyms.values() for s in v), e.description]).lower()

    return [
        e
        for e in db.query(GlossaryEntry).order_by(GlossaryEntry.id.desc())
        if (not status or e.status == status) and (not category or e.category == category) and q.lower() in text(e)
    ]


@router.post("", response_model=GlossaryOut, summary="Create an entry (translators create drafts)")
def create_entry(db: DB, user: Translator, body: GlossaryIn) -> GlossaryEntry:
    terms, syn = clean(body)
    if (dups := duplicates(db, terms)) and not body.force:
        raise conflict(dups)
    status = "draft" if user.role == "translator" else body.status or "approved"
    e = GlossaryEntry(
        terms=terms,
        synonyms=syn,
        description=body.description,
        category=body.category,
        do_not_translate=body.do_not_translate,
        status=status,
        updated_by=user.username,
    )
    db.add(e)
    db.flush()
    snapshot(db, e, "create", user.username)
    db.commit()
    return e


@router.put("/{eid}", response_model=GlossaryOut)
def update_entry(db: DB, user: Reviewer, eid: int, body: GlossaryIn) -> GlossaryEntry:
    e = find(db, eid)
    terms, syn = clean(body)
    if (dups := duplicates(db, terms, skip=eid)) and not body.force:
        raise conflict(dups)
    e.terms, e.synonyms, e.description, e.category = terms, syn, body.description, body.category
    e.do_not_translate, e.status = body.do_not_translate, body.status or e.status
    e.version, e.updated_by = e.version + 1, user.username
    snapshot(db, e, "update", user.username)
    db.commit()
    return e


@router.post("/{eid}/status", response_model=GlossaryOut, summary="Approve or reject an entry")
def set_status(db: DB, user: Reviewer, eid: int, status: str) -> GlossaryEntry:
    if status not in ("draft", "approved", "rejected"):
        raise HTTPException(400, "Ungültiger Status")
    e = find(db, eid)
    e.status, e.version, e.updated_by = status, e.version + 1, user.username
    snapshot(db, e, "update", user.username)
    db.commit()
    return e


@router.delete("/{eid}")
def delete_entry(db: DB, user: Admin, eid: int) -> dict[str, bool]:
    e = find(db, eid)
    e.version += 1
    snapshot(db, e, "delete", user.username)
    db.delete(e)
    db.commit()
    return {"ok": True}


@router.get("/{eid}/versions", response_model=list[GlossaryVersionOut], summary="Version history of an entry")
def versions(db: DB, _: Me, eid: int) -> list[GlossaryVersion]:
    return db.query(GlossaryVersion).filter(GlossaryVersion.entry_id == eid).order_by(GlossaryVersion.id.desc()).all()


@router.get("/duplicates/scan", response_model=list[DuplicateGroup], summary="Find duplicate terms")
def scan_duplicates(db: DB, _: Me) -> list[DuplicateGroup]:
    seen: dict[tuple[str, str], int] = {}
    groups: list[DuplicateGroup] = []
    for e in db.query(GlossaryEntry):
        for lang, t in e.terms.items():
            if seen.setdefault((lang, norm(t)), e.id) != e.id:
                groups.append(DuplicateGroup(term=t, lang=lang, ids=[seen[(lang, norm(t))], e.id]))
    return groups


@router.get("/export/xlsx", response_class=StreamingResponse, summary="Export the glossary to Excel")
def export_xlsx(db: DB, _: Me) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Glossar"
    ws.append(COLUMNS)
    for e in db.query(GlossaryEntry).order_by(GlossaryEntry.id):
        ws.append(
            [
                e.terms.get("de", ""),
                e.terms.get("en", ""),
                "; ".join(e.synonyms.get("de", [])),
                "; ".join(e.synonyms.get("en", [])),
                e.description,
                "ja" if e.do_not_translate else "",
                e.category,
                e.status,
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type=XLSX, headers={"Content-Disposition": 'attachment; filename="glossar.xlsx"'}
    )


@router.post("/import/xlsx", response_model=ImportResult, summary="Import glossary entries from Excel")
def import_xlsx(
    db: DB, user: Reviewer, file: Annotated[UploadFile, File()], update_existing: bool = False
) -> ImportResult:
    """Columns: de, en, de_synonyme, en_synonyme, beschreibung, nicht_uebersetzen, kategorie, status."""
    try:
        rows = list(load_workbook(file.file, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    except Exception:
        raise HTTPException(400, "Excel-Datei konnte nicht gelesen werden.") from None
    header = [ALIASES.get(str(h or "").strip().lower(), str(h or "").strip().lower()) for h in rows[0]] if rows else []
    if "de" not in header and "en" not in header:
        raise HTTPException(400, "Spalten „de“ und/oder „en“ fehlen in der Kopfzeile.")
    res = ImportResult(created=0, updated=0, skipped=0, skipped_rows=[])
    for n, row in enumerate(rows[1:], start=2):
        r = {h: str(v).strip() for h, v in zip(header, row, strict=False) if v is not None}
        terms = {lang: r[lang] for lang in LANGS if r.get(lang)}
        dnt = r.get("nicht_uebersetzen", "").lower() in ("ja", "yes", "true", "1", "x")
        syn = {
            lang: [s.strip() for s in re.split(r"[;|]", r.get(f"{lang}_synonyme", "")) if s.strip()] for lang in LANGS
        }
        dups = duplicates(db, terms) if terms else []
        reason = (
            "Begriff oder Übersetzung fehlt"
            if not terms or (not dnt and len(terms) < 2)
            else f"Dublette: {next(iter(terms.values()))}"
            if dups and not update_existing
            else ""
        )
        if reason:
            res.skipped += 1
            res.skipped_rows.append({"row": n, "reason": reason})
            continue
        if dups:
            e = dups[0]
            e.terms, e.synonyms, e.do_not_translate = {**e.terms, **terms}, syn, dnt
            e.description, e.version, e.updated_by = (
                r.get("beschreibung") or e.description,
                e.version + 1,
                user.username,
            )
            res.updated += 1
        else:
            status = r.get("status") if r.get("status") in ("draft", "approved") else "approved"
            e = GlossaryEntry(
                terms=terms,
                synonyms=syn,
                description=r.get("beschreibung", ""),
                do_not_translate=dnt,
                category=r.get("kategorie") or "technical",
                status=status or "approved",
                updated_by=user.username,
            )
            db.add(e)
            db.flush()
            res.created += 1
        snapshot(db, e, "update" if dups else "create", user.username)
    db.commit()
    return res
